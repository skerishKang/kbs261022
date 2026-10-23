#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_full_catalog.py -- CGI 전체 상품 DB 통합 빌더 (이슈 #19).

역할
    공급처 공개 수집 원본(2,281 + 2,294 = 4,575건)을 **이슈 #8 의 통합 상품
    마스터 스키마**로 정규화하고, 그 결과로 다음을 만든다.

      1) product_master.full.json / .csv   PR #16 스키마 마스터 (4,575행)
      2) category_catalog.csv              카테고리 카탈로그
      3) product_db.xlsx                   필수 7개 시트 + 옵션 SKU 시트
      4) catalog_source.csv                PR #17 카탈로그 빌더가 소비하는 형식
      5) verification.json                 검증 결과

재사용
    * 스키마·정규화·중복 판정: tools/product_master/ (이슈 #8, PR #16)
    * 수집 원본/파서: tools/full_catalog/collect_cgi_catalog.py (PR #12 Model B)
    * 카탈로그 사이트 JSON: tools/build_catalog_data.py (PR #17) 를 별도 호출

원칙
    * 원천에 없는 값은 만들지 않는다. 채워지지 않은 필드는 그대로 공란.
    * 가격·재고·도매가는 공개 표기만 기록한다.
    * 사이트별 cate_no 는 namespace 로 분리된 상태로만 비교한다.
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from tools.product_master import dedupe as DD  # noqa: E402
from tools.product_master import normalize as N  # noqa: E402
from tools.product_master import schema as SC  # noqa: E402
from tools.product_master import validate as VA  # noqa: E402

#: 수집 원본의 사이트 → (표시명, 이미지 절대화 호스트, namespace)
SITE_INFO: Dict[str, Tuple[str, str, str]] = {
    "samasound": ("삼아사운드", "www.samasound.co.kr", SC.SITES["samasound"]["namespace"]),
    "samamusic": ("삼아스토어", "www.samamusic.co.kr", SC.SITES["samamusic"]["namespace"]),
}

#: 수집 상태 → PR #16 ``price_type``
PRICE_STATE_MAP: Dict[str, str] = {
    "numeric_price": "public_selling_price",
    "inquiry_price": "inquiry",
    "sold_out_sale_completed": "unknown",
    "other": "unknown",
}

#: ``canonical`` URL 호스트 (수집 호스트와 다를 수 있다)
CANONICAL_HOSTS = {
    "samasound": "samasound.co.kr",
    "samamusic": "samastore.co.kr",
}


# --------------------------------------------------------------------------
# 유틸
# --------------------------------------------------------------------------
def now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def parse_json_field(raw: Any, default: Any) -> Any:
    if raw in (None, ""):
        return default
    if isinstance(raw, (list, dict)):
        return raw
    try:
        return json.loads(raw)
    except (ValueError, TypeError):
        return default


def absolute_image(url: str, site: str) -> str:
    """상대 이미지 경로를 원본 도메인으로 절대화한다."""
    if not url:
        return ""
    url = url.strip()
    if url.startswith("//"):
        return "https:" + url
    if url.startswith("/"):
        host = SITE_INFO.get(site, ("", "www.example.co.kr", ""))[1]
        return f"https://{host}{url}"
    return url


def split_images(raw: Any, site: str) -> List[str]:
    """'|' 구분 / JSON 배열 / 단일 URL 을 모두 처리."""
    if not raw:
        return []
    if isinstance(raw, list):
        candidates = [str(x).strip() for x in raw]
    else:
        text = str(raw).strip()
        if not text:
            return []
        if text.startswith("["):
            parsed = parse_json_field(text, [])
            candidates = [str(x).strip() for x in parsed] if isinstance(parsed, list) else []
        else:
            candidates = [p.strip() for p in text.split("|") if p.strip()]
    out: List[str] = []
    for url in candidates:
        full = absolute_image(url, site)
        if full.startswith("http") and full not in out:
            out.append(full)
    return out


def split_path(raw: str) -> List[str]:
    """'홈 > 제품분류별 보기 > 마이크' -> ['홈','제품분류별 보기','마이크']"""
    if not raw:
        return []
    out = []
    for seg in str(raw).replace("|", ">").split(">"):
        seg = " ".join(seg.split()).strip(" -·")
        if seg:
            out.append(seg)
    return out


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


# --------------------------------------------------------------------------
# 1. 수집 원본 → PR #16 마스터 레코드
# --------------------------------------------------------------------------
def to_master_record(raw: Mapping[str, Any], extraction_time: str) -> Optional[Dict[str, Any]]:
    """수집 원본 1건 → PR #16 스키마 마스터 레코드 1건."""
    site = str(raw.get("source_site") or "").strip()
    if site not in SITE_INFO:
        return None
    label, _, namespace = SITE_INFO[site]

    product_id = str(raw.get("source_product_id") or "").strip()
    name = N.normalize_name(raw.get("product_name"))
    if not product_id or not name:
        return None

    specs = parse_json_field(raw.get("specs"), {})
    if not isinstance(specs, dict):
        specs = {}

    # --- 브랜드 / 제조사 / 모델 (출처를 함께 기록) --------------------------
    spec_brand = N.normalize_brand(specs.get("브랜드")) if specs.get("브랜드") else None
    spec_manufacturer = N.normalize_manufacturer(specs.get("제조사")) if specs.get("제조사") else None
    spec_model = N.normalize_name(specs.get("모델명")) if specs.get("모델명") else None

    # samasound 는 브랜드 스펙행이 없다. 대신 상품이 게시된 브랜드 카테고리에서
    # 파생된 값이 원본 증거로 존재한다. 근거 수준을 brand_source 로 명시한다.
    derived_brand = N.normalize_brand(raw.get("source_brand")) if raw.get("source_brand") else None
    if spec_brand:
        brand, brand_source = spec_brand, "spec"
    elif derived_brand:
        brand, brand_source = derived_brand, "category"
    else:
        brand, brand_source = None, "none"

    if spec_manufacturer:
        manufacturer, manufacturer_source = spec_manufacturer, "spec"
    else:
        manufacturer, manufacturer_source = None, "none"

    model = spec_model  # 모델명 스펙행이 없으면 null (이름 파싱으로 만들지 않는다)

    # --- 가격 ---------------------------------------------------------------
    price_state = str(raw.get("published_price_state") or "").strip()
    price_text = N.normalize_text(raw.get("published_price_text")) or None
    price_value = raw.get("published_price")
    price_type = PRICE_STATE_MAP.get(price_state, "unknown")
    if price_type == "public_selling_price":
        try:
            public_price = int(str(price_value).replace(",", "").strip())
        except (TypeError, ValueError):
            public_price, price_type = None, "unknown"
    else:
        public_price = None
    if public_price is None and price_type == "public_selling_price":
        price_type = "unknown"

    # --- 카테고리 -----------------------------------------------------------
    cate_id = str(raw.get("source_category_id") or "").strip() or None
    cat_name = N.normalize_text(raw.get("source_category_name")) or None
    category_path = split_path(raw.get("source_category_path") or "")
    if not category_path and cat_name:
        category_path = [cat_name]

    # --- 옵션 ---------------------------------------------------------------
    options: List[Dict[str, Any]] = []
    summary = parse_json_field(raw.get("option_summary"), {})
    if isinstance(summary, dict):
        for group in ("main", "addproduct_options", "add_option_data", "addproducts"):
            for entry in summary.get(group) or []:
                if not isinstance(entry, dict):
                    continue
                values = []
                codes = entry.get("item_codes") or []
                for idx, value in enumerate(entry.get("values") or []):
                    label = N.normalize_text(value)
                    if not label:
                        continue
                    delta = None
                    m = re.search(r"\(([+-]?[0-9][0-9,]*)\s*원\)", label)
                    values.append(
                        {
                            "label": label,
                            "price_delta": delta,
                            "soldout": "[품절]" in label or "품절" in label,
                            "item_code": codes[idx] if idx < len(codes) else None,
                        }
                    )
                if not values:
                    continue
                options.append(
                    {
                        "label": N.normalize_text(entry.get("option_name")) or "옵션",
                        "group": group,
                        "values": values,
                    }
                )

    # --- 파생 매칭 키 -------------------------------------------------------
    phrases = N.model_phrases(name, brand=brand, source_model=model)
    derived = {
        "model_phrases": phrases,
        "model_phrase_compact": sorted({c for c in (N.model_compact(p) for p in phrases) if c}),
        "variant_signature": sorted(N.variant_signature(name)),
        "derivation_note": (
            "파생 매칭 전용 값. 원천 'model' 필드와 별개이며 모델명 스펙행이 "
            "없는 사이트는 model=null 로 유지된다."
        ),
    }

    # --- 품질 플래그 --------------------------------------------------------
    flags: List[str] = []
    if public_price is None:
        flags.append("public_price_missing")
    if price_state == "sold_out_sale_completed":
        flags.append("sold_out")
    if price_state == "inquiry_price":
        flags.append("public_price_inquiry_only")
    if not split_images(raw.get("image_urls"), site):
        flags.append("image_missing")
    if brand_source == "category":
        flags.append("brand_from_category_not_spec")
    if brand is None:
        flags.append("brand_missing")
    if manufacturer is None:
        flags.append("manufacturer_missing")
    if model is None:
        flags.append("model_missing")
    if not options:
        flags.append("options_missing")
    if not category_path:
        flags.append("category_path_missing")
    if price_state == "other":
        flags.append("price_text_unclassified")

    evidence = {
        "url": raw.get("canonical_url"),
        "list_url": raw.get("source_list_url"),
        "source_site_label": raw.get("source_site_label") or label,
        "observed_price_text": price_text,
        "observed_price_state": price_state,
        "observed_brand": raw.get("source_brand"),
        "observed_manufacturer": raw.get("manufacturer"),
        "observed_model_spec": spec_model,
        "spec_row_labels": sorted(specs.keys()),
        "notes": "공급처 공개 HTML 수집 원본 (issues #19 전수 수집)",
    }

    return {
        "source_site": site,
        "source_product_id": product_id,
        "source_key": SC.source_key(site, product_id),
        "canonical_url": raw.get("canonical_url"),
        "collection_url": raw.get("source_list_url"),
        "name": name,
        "brand": brand,
        "brand_source": brand_source,
        "manufacturer": manufacturer,
        "manufacturer_source": manufacturer_source,
        "model": model,
        "category_path": category_path,
        "category_namespace": namespace,
        "category_path_cate_no": [],
        "category_leaf_cate_no": cate_id,
        "specs": specs,
        "options": options,
        "image_urls": split_images(raw.get("image_urls"), site),
        "public_price": public_price,
        "public_price_text": price_text,
        "price_type": price_type,
        "dealer_price": None,
        "availability": None,
        "extraction_time": raw.get("updated_at") or extraction_time,
        "extraction_method": raw.get("extraction_method") or "http_get_public_html",
        "source_evidence": evidence,
        "verification_status": "http_verified",
        "derived": derived,
        "naver_product_id": None,
        "erp_item_code": None,
        "smartstore_category_no": None,
        "mapping_ready": False,
        "quality_flags": flags,
        "dedupe_group_id": None,
        "dedupe_confidence": None,
        # 공급처 원본 식별자 (마스터 밖에서 보존)
        "source_item_code": raw.get("source_item_code"),
        "product_code": raw.get("product_code"),
        "source_category_name": cat_name,
    }


# --------------------------------------------------------------------------
# 2. 시트 빌더
# --------------------------------------------------------------------------
def build_unified_sheet(master: Sequence[Mapping[str, Any]]) -> List[Dict[str, Any]]:
    """통합 상품 마스터 시트."""
    rows = []
    for r in master:
        rows.append(
            {
                "통합키": r["source_key"],
                "공급처": SITE_INFO[r["source_site"]][0],
                "공급처코드": r["source_site"],
                "공급처원본ID": r["source_product_id"],
                "상품명": r["name"],
                "브랜드": r["brand"] or "",
                "브랜드근거": r["brand_source"],
                "제조사": r["manufacturer"] or "",
                "제조사근거": r["manufacturer_source"],
                "모델명": r["model"] or "",
                "카테고리경로": " > ".join(r["category_path"]),
                "카테고리namespace": r["category_namespace"],
                "말단cate_no": r["category_leaf_cate_no"] or "",
                "공개판매가": r["public_price"] if r["public_price"] is not None else "",
                "가격표시": r["public_price_text"] or "",
                "가격유형": r["price_type"],
                "도매가": "",
                "재고": "",
                "옵션수": sum(len(o["values"]) for o in r["options"]),
                "이미지수": len(r["image_urls"]),
                "대표이미지": r["image_urls"][0] if r["image_urls"] else "",
                "규격행수": len(r["specs"]),
                "출처URL": r["canonical_url"] or "",
                "목록URL": r["collection_url"] or "",
                "수집시각": r["extraction_time"],
                "수집방식": r["extraction_method"],
                "검증상태": r["verification_status"],
                "중복그룹": r["dedupe_group_id"] or "",
                "중복판정": r["dedupe_confidence"] or "",
                "품질플래그": "|".join(r["quality_flags"]),
                "네이버상품ID": "",
                "ERP품목코드": "",
                "스마트스토어카테고리": "",
                "매핑준비": "아니오",
            }
        )
    return rows


def build_supplier_sheet(
    master: Sequence[Mapping[str, Any]],
    raw_by_key: Mapping[str, Mapping[str, Any]],
) -> List[Dict[str, Any]]:
    """공급처별 원본 상품 시트 — 원천 식별자와 출처를 그대로 보존."""
    rows = []
    for r in master:
        raw = raw_by_key.get(r["source_key"], {})
        rows.append(
            {
                "공급처": SITE_INFO[r["source_site"]][0],
                "공급처코드": r["source_site"],
                "공급처원본ID": r["source_product_id"],
                "상품코드": raw.get("product_code") or "",
                "품목코드": raw.get("source_item_code") or "",
                "원본상품명": raw.get("product_name") or r["name"],
                "원본브랜드": raw.get("source_brand") or "",
                "원본제조사": raw.get("manufacturer") or "",
                "원본모델명": raw.get("model") or "",
                "원본카테고리ID": raw.get("source_category_id") or "",
                "원본카테고리명": raw.get("source_category_name") or "",
                "원본카테고리경로": raw.get("source_category_path") or "",
                "원본가격표시": raw.get("published_price_text") or "",
                "원본가격상태": raw.get("published_price_state") or "",
                "상세수집여부": "예" if raw.get("detail_collected") in (True, "1", "True", "true") else "아니오",
                "canonical_url": raw.get("canonical_url") or "",
                "목록URL": raw.get("source_list_url") or "",
                "수집시각": raw.get("updated_at") or r["extraction_time"],
                "검증상태": raw.get("verification_status") or "",
                "출처증거": raw.get("source_evidence") or "",
            }
        )
    return rows


def build_brand_category_sheet(master: Sequence[Mapping[str, Any]]) -> List[Dict[str, Any]]:
    """브랜드·카테고리 시트 — 사이트 namespace 를 반드시 함께 표기."""
    agg: Dict[Tuple[str, str, str, str], int] = Counter()
    for r in master:
        brand = r["brand"] or "(미확인)"
        cat = r["category_leaf_cate_no"] or "(미확인)"
        cat_name = r.get("source_category_name") or ""
        agg[(r["source_site"], r["category_namespace"], brand, f"{cat} {cat_name}".strip())] += 1
    rows = []
    for (site, namespace, brand, cat), count in sorted(
        agg.items(), key=lambda kv: (kv[0][0], -kv[1], kv[0][2])
    ):
        rows.append(
            {
                "공급처코드": site,
                "cate_no_namespace": namespace,
                "브랜드": brand,
                "브랜드근거": "spec"
                if namespace == "samamusic.cate_no"
                else "category",
                "cate_no": cat.split(" ")[0] if cat != "(미확인)" else "",
                "카테고리명": " ".join(cat.split(" ")[1:]) if cat != "(미확인)" else "",
                "상품수": count,
            }
        )
    return rows


def build_duplicate_sheet(
    candidates: Sequence[DD.MatchCandidate],
    master: Sequence[Mapping[str, Any]],
) -> List[Dict[str, Any]]:
    """중복 확정 및 검토 후보 시트."""
    by_key = {r["source_key"]: r for r in master}
    rows = []
    for c in candidates:
        left, right = by_key.get(c.left["source_key"]), by_key.get(c.right["source_key"])
        if not left or not right:
            continue
        rows.append(
            {
                "판정": c.verdict,
                "자동병합가능": "예" if c.auto_mergeable else "아니오",
                "왼쪽공급처": SITE_INFO[left["source_site"]][0],
                "왼쪽원본ID": left["source_product_id"],
                "왼쪽상품명": left["name"],
                "왼쪽브랜드": left["brand"] or "",
                "왼쪽모델명": left["model"] or "",
                "왼쪽가격": left["public_price"] if left["public_price"] is not None else "",
                "오른쪽공급처": SITE_INFO[right["source_site"]][0],
                "오른쪽원본ID": right["source_product_id"],
                "오른쪽상품명": right["name"],
                "오른쪽브랜드": right["brand"] or "",
                "오른쪽모델명": right["model"] or "",
                "오른쪽가격": right["public_price"] if right["public_price"] is not None else "",
                "공통모델표현": "|".join(c.shared_phrases),
                "이름유사도": c.signals.get("name_similarity"),
                "변형충돌": "예" if c.signals.get("variant_conflict") else "아니오",
                "브랜드값일치": "예" if c.signals.get("org_value_agreement") else "아니오",
                "사유": " / ".join(c.reasons),
            }
        )
    return rows


def build_gap_sheet(master: Sequence[Mapping[str, Any]]) -> List[Dict[str, Any]]:
    """가격·이미지·규격 누락 현황 시트."""
    rows = []
    for r in master:
        rows.append(
            {
                "통합키": r["source_key"],
                "공급처": SITE_INFO[r["source_site"]][0],
                "공급처원본ID": r["source_product_id"],
                "상품명": r["name"],
                "공개판매가": r["public_price"] if r["public_price"] is not None else "",
                "가격표시": r["public_price_text"] or "",
                "가격유형": r["price_type"],
                "가격누락": "예" if r["public_price"] is None else "아니오",
                "도매가": "",
                "재고": "",
                "이미지수": len(r["image_urls"]),
                "이미지누락": "예" if not r["image_urls"] else "아니오",
                "규격행수": len(r["specs"]),
                "규격누락": "예" if not r["specs"] else "아니오",
                "브랜드": r["brand"] or "",
                "브랜드근거": r["brand_source"],
                "제조사": r["manufacturer"] or "",
                "모델명": r["model"] or "",
                "옵션수": sum(len(o["values"]) for o in r["options"]),
                "누락플래그": "|".join(r["quality_flags"]),
            }
        )
    return rows


def build_naver_sheet(master: Sequence[Mapping[str, Any]]) -> List[Dict[str, Any]]:
    """네이버 스마트스토어 등록 준비 시트 — 아직 ID 는 전부 공란."""
    rows = []
    for r in master:
        price_ok = r["public_price"] is not None
        rows.append(
            {
                "공급처원본ID": f"{r['source_site']}:{r['source_product_id']}",
                "상품명": r["name"],
                "브랜드": r["brand"] or "",
                "모델명": r["model"] or "",
                "카테고리경로": " > ".join(r["category_path"]),
                "판매가": r["public_price"] if r["public_price"] is not None else "",
                "가격확인필요": "아니오" if price_ok else "예",
                "대표이미지": r["image_urls"][0] if r["image_urls"] else "",
                "상세이미지수": max(0, len(r["image_urls"]) - 1),
                "설명": r["specs"].get("상품요약정보", "") or r["specs"].get("상품간략설명", ""),
                "옵션수": sum(len(o["values"]) for o in r["options"]),
                "네이버상품ID": "",
                "스마트스토어카테고리번호": "",
                "등록상태": "미등록",
                "등록가능여부": "예" if (price_ok and r["image_urls"]) else "아니오",
                "사유": "가격 문의" if not price_ok else ("이미지 없음" if not r["image_urls"] else ""),
            }
        )
    return rows


def build_erp_sheet(master: Sequence[Mapping[str, Any]]) -> List[Dict[str, Any]]:
    """이카운트 품목 매핑 준비 시트 — 아직 ERP 코드는 전부 공란."""
    rows = []
    for r in master:
        rows.append(
            {
                "공급처원본ID": f"{r['source_site']}:{r['source_product_id']}",
                "상품명": r["name"],
                "브랜드": r["brand"] or "",
                "제조사": r["manufacturer"] or "",
                "모델명": r["model"] or "",
                "공급처": SITE_INFO[r["source_site"]][0],
                "원본품목코드": r.get("source_item_code") or "",
                "원본상품코드": r.get("product_code") or "",
                "옵션SKU수": sum(len(o["values"]) for o in r["options"]),
                "규격행수": len(r["specs"]),
                "이카운트품목코드": "",
                "이카운트거래처코드": "",
                "ERP매핑상태": "미매핑",
                "비고": "스토어 개설 전이므로 ERP 코드를 생성하지 않음",
            }
        )
    return rows


def build_option_sheet(master: Sequence[Mapping[str, Any]]) -> List[Dict[str, Any]]:
    """옵션 SKU 시트 (보조)."""
    rows = []
    for r in master:
        for opt in r["options"]:
            for value in opt["values"]:
                rows.append(
                    {
                        "공급처원본ID": f"{r['source_site']}:{r['source_product_id']}",
                        "상품명": r["name"],
                        "옵션명": opt["label"],
                        "옵션값": value["label"],
                        "추가금": value["price_delta"] if value["price_delta"] is not None else "",
                        "품절": "예" if value["soldout"] else "아니오",
                        "품목코드": value.get("item_code") or "",
                        "옵션그룹": opt["group"],
                    }
                )
    return rows


# --------------------------------------------------------------------------
# 3. 카탈로그 빌더 입력 (PR #17 형식)
# --------------------------------------------------------------------------
def build_catalog_source_csv(
    master: Sequence[Mapping[str, Any]], raw_by_key: Mapping[str, Mapping[str, Any]], path: Path
) -> int:
    """PR #17 의 tools/build_catalog_data.py 가 그대로 읽을 수 있는 CSV."""
    fields = [
        "source_site",
        "source_site_label",
        "source_product_id",
        "source_item_code",
        "product_code",
        "product_name",
        "source_brand",
        "manufacturer",
        "model",
        "source_category_id",
        "source_category_name",
        "source_category_path",
        "published_price",
        "published_price_text",
        "specs",
        "source_evidence",
        "image_urls",
        "option_summary",
        "availability_unverified",
        "canonical_url",
        "source_list_url",
        "extraction_method",
        "verification_status",
        "updated_at",
    ]
    count = 0
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        for r in master:
            raw = raw_by_key.get(r["source_key"], {})
            writer.writerow(
                {
                    "source_site": r["source_site"],
                    "source_site_label": SITE_INFO[r["source_site"]][0],
                    "source_product_id": r["source_product_id"],
                    "source_item_code": raw.get("source_item_code") or "",
                    "product_code": raw.get("product_code") or "",
                    "product_name": r["name"],
                    "source_brand": r["brand"] or "",
                    "manufacturer": r["manufacturer"] or "",
                    "model": r["model"] or "",
                    "source_category_id": r["category_leaf_cate_no"] or "",
                    "source_category_name": r.get("source_category_name") or "",
                    "source_category_path": " > ".join(r["category_path"]),
                    "published_price": r["public_price"] if r["public_price"] is not None else "",
                    "published_price_text": r["public_price_text"] or "",
                    "specs": json.dumps(r["specs"], ensure_ascii=False),
                    "source_evidence": json.dumps(r["source_evidence"], ensure_ascii=False),
                    "image_urls": "|".join(r["image_urls"]),
                    "option_summary": json.dumps(
                        {
                            "main": [
                                {
                                    "option_name": o["label"],
                                    "values": [v["label"] for v in o["values"]],
                                    "item_codes": [v.get("item_code") for v in o["values"]],
                                }
                                for o in r["options"]
                                if o["group"] == "main"
                            ],
                            "addproduct_options": [],
                            "addproducts": [],
                            "add_option_data": [],
                        },
                        ensure_ascii=False,
                    ),
                    "availability_unverified": raw.get("availability_unverified") or "",
                    "canonical_url": r["canonical_url"] or "",
                    "source_list_url": r["collection_url"] or "",
                    "extraction_method": r["extraction_method"],
                    "verification_status": r["verification_status"],
                    "updated_at": r["extraction_time"],
                }
            )
            count += 1
    return count


def build_category_catalog_csv(
    master: Sequence[Mapping[str, Any]], path: Path
) -> int:
    """사이트별 카테고리 카탈로그 (cate_no namespace 분리)."""
    agg: Dict[Tuple[str, str, str, str], Dict[str, Any]] = {}
    for r in master:
        cid = r["category_leaf_cate_no"] or ""
        if not cid:
            continue
        key = (r["source_site"], cid)
        entry = agg.setdefault(
            key,
            {
                "source_site": r["source_site"],
                "category_id": cid,
                "category_name": r.get("source_category_name") or "",
                "category_path": " > ".join(r["category_path"]),
                "category_level": str(len(r["category_path"])),
                "parent_category_id": "",
                "registered_count": 0,
                "list_url": "",
            },
        )
        entry["registered_count"] += 1
    fields = [
        "source_site",
        "category_id",
        "category_name",
        "category_path",
        "category_level",
        "parent_category_id",
        "registered_count",
        "list_url",
    ]
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        for (site, cid), entry in sorted(agg.items()):
            host = CANONICAL_HOSTS[site]
            entry["list_url"] = f"https://{host}/product/list.html?cate_no={cid}"
            writer.writerow(entry)
    return len(agg)


# --------------------------------------------------------------------------
# 4. XLSX
# --------------------------------------------------------------------------
SHEET_ORDER: List[Tuple[str, str]] = [
    ("통합상품", "통합 상품 마스터"),
    ("공급처원본", "공급처별 원본 상품"),
    ("브랜드_카테고리", "브랜드·카테고리"),
    ("중복후보", "중복 확정 및 검토 후보"),
    ("누락_검증현황", "가격·이미지·규격 누락 현황"),
    ("네이버스토어준비", "네이버 스마트스토어 등록 준비"),
    ("이카운트품목매핑", "이카운트 품목 매핑"),
    ("옵션SKU", "옵션 SKU (보조)"),
]


def write_xlsx(path: Path, sheets: List[Tuple[str, List[Dict[str, Any]]]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F3A93")

    for name, rows in sheets:
        ws = wb.create_sheet(title=name[:31])
        if not rows:
            ws.append(["(데이터 없음)"])
            continue
        fields = list(rows[0].keys())
        ws.append(fields)
        for c in ws[1]:
            c.font = head_font
            c.fill = head_fill
            c.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append([r.get(f, "") for f in fields])
        ws.freeze_panes = "A2"
        for i, f in enumerate(fields, start=1):
            width = max(len(str(f)) + 2, 12)
            ws.column_dimensions[get_column_letter(i)].width = min(width, 46)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# --------------------------------------------------------------------------
# 5. 검증
# --------------------------------------------------------------------------
def verify(
    master: Sequence[Mapping[str, Any]],
    candidates: Sequence[DD.MatchCandidate],
    groups: Mapping[str, Sequence[str]],
    raw_counts: Mapping[str, int],
    site_json_count: Optional[int],
) -> Dict[str, Any]:
    by_site = Counter(r["source_site"] for r in master)
    schema_result = VA.validate_dataset(master)
    same_site = DD.find_same_site_duplicates(master)

    verdict_counts = Counter(c.verdict for c in candidates)
    missing_price = sum(1 for r in master if r["public_price"] is None)
    missing_image = sum(1 for r in master if not r["image_urls"])
    missing_specs = sum(1 for r in master if not r["specs"])
    missing_brand = sum(1 for r in master if not r["brand"])
    missing_model = sum(1 for r in master if not r["model"])
    missing_manufacturer = sum(1 for r in master if not r["manufacturer"])

    return {
        "generated_at": now_iso(),
        "source_products": {
            "samasound": raw_counts.get("samasound", 0),
            "samamusic": raw_counts.get("samamusic", 0),
            "total": sum(raw_counts.values()),
        },
        "master_products": len(master),
        "master_products_by_site": dict(by_site),
        "source_count_match": all(
            by_site.get(site, 0) == count for site, count in raw_counts.items()
        ),
        "schema": {
            "ok": schema_result["ok"],
            "structural_error_count": schema_result["structural_error_count"],
            "policy_error_count": schema_result["policy_error_count"],
            "unique_key_count": schema_result["unique_key_count"],
        },
        "same_site_duplicate_keys": len(same_site),
        "duplicates": {
            "candidates": len(candidates),
            "verdict_counts": dict(verdict_counts),
            "confirmed": verdict_counts.get("confirmed", 0),
            "candidate": verdict_counts.get("candidate", 0),
            "probable": verdict_counts.get("probable", 0),
            "auto_merge_groups": len(groups),
        },
        "missing_rates": {
            "public_price": {"missing": missing_price, "rate": round(missing_price / len(master), 4) if master else 0},
            "image": {"missing": missing_image, "rate": round(missing_image / len(master), 4) if master else 0},
            "specs": {"missing": missing_specs, "rate": round(missing_specs / len(master), 4) if master else 0},
            "brand": {"missing": missing_brand, "rate": round(missing_brand / len(master), 4) if master else 0},
            "model": {"missing": missing_model, "rate": round(missing_model / len(master), 4) if master else 0},
            "manufacturer": {"missing": missing_manufacturer, "rate": round(missing_manufacturer / len(master), 4) if master else 0},
        },
        "category_namespace": {
            "samasound": SC.SITES["samasound"]["namespace"],
            "samamusic": SC.SITES["samamusic"]["namespace"],
            "distinct_leaf_cate_no": {
                site: len({r["category_leaf_cate_no"] for r in master if r["source_site"] == site})
                for site in sorted(by_site)
            },
            "collision_note": "cate_no 42 는 samasound=Shure, samamusic=스피커 로 서로 다른 의미",
        },
        "site_catalog_json": {
            "products": site_json_count,
            "matches_master": site_json_count == len(master) if site_json_count is not None else None,
        },
    }


# --------------------------------------------------------------------------
# main
# --------------------------------------------------------------------------
def main() -> int:
    ap = argparse.ArgumentParser(description="CGI 전체 상품 DB 통합 빌더 (이슈 #19)")
    ap.add_argument("--data-dir", default="C:/cgi-work/catalog-data")
    ap.add_argument("--out-dir", default="data/full_catalog")
    ap.add_argument("--xlsx", default=None, help="기본 <out-dir>/product_db.xlsx")
    ap.add_argument("--site-json", default=None, help="PR #17 빌더가 만든 products.json (검증용)")
    ap.add_argument("--no-xlsx", action="store_true")
    args = ap.parse_args()

    data_dir = Path(args.data_dir)
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx_path = Path(args.xlsx) if args.xlsx else out_dir / "product_db.xlsx"

    extraction_time = now_iso()
    raw_by_key: Dict[str, Mapping[str, Any]] = {}
    raw_counts: Dict[str, int] = {}
    master: List[Dict[str, Any]] = []

    for site in ("samasound", "samamusic"):
        path = data_dir / f"{site}_detail.json"
        if not path.exists():
            print(f"[WARN] 원본 없음: {path}")
            continue
        rows = load_json(path)
        raw_counts[site] = len(rows)
        for raw in rows:
            record = to_master_record(raw, extraction_time)
            if record is None:
                continue
            master.append(record)
            raw_by_key[record["source_key"]] = raw

    print(f"[1/6] 마스터 레코드 생성: {len(master)}건 ({raw_counts})")

    # --- 중복 판정 (PR #16 알고리즘) ---------------------------------------
    candidates = DD.find_candidates(master)
    groups = DD.build_groups(master, candidates)
    master = DD.apply_groups(master, groups)
    by_key = {r["source_key"]: r for r in master}
    counts = Counter(c.verdict for c in candidates)
    print(
        f"[2/6] 중복 판정: 후보 {len(candidates)} (confirmed {counts.get('confirmed', 0)}, "
        f"candidate {counts.get('candidate', 0)}, probable {counts.get('probable', 0)}) "
        f"자동 병합 그룹 {len(groups)}"
    )

    # --- 산출물 -------------------------------------------------------------
    SC.write_json(master, str(out_dir / "product_master.full.json"))
    SC.write_csv(master, str(out_dir / "product_master.full.csv"))
    print(f"[3/6] 마스터 파일: {out_dir / 'product_master.full.json'}")

    build_catalog_source_csv(master, raw_by_key, out_dir / "catalog_source.csv")
    build_category_catalog_csv(master, out_dir / "category_catalog.csv")
    print(f"[4/6] 카탈로그 입력 CSV 생성")

    sheets = [
        ("통합상품", build_unified_sheet(master)),
        ("공급처원본", build_supplier_sheet(master, raw_by_key)),
        ("브랜드_카테고리", build_brand_category_sheet(master)),
        ("중복후보", build_duplicate_sheet(candidates, master)),
        ("누락_검증현황", build_gap_sheet(master)),
        ("네이버스토어준비", build_naver_sheet(master)),
        ("이카운트품목매핑", build_erp_sheet(master)),
        ("옵션SKU", build_option_sheet(master)),
    ]
    if not args.no_xlsx:
        write_xlsx(xlsx_path, sheets)
        print(f"[5/6] XLSX: {xlsx_path} ({xlsx_path.stat().st_size} bytes)")

    site_json_count = None
    if args.site_json and Path(args.site_json).exists():
        site_json_count = len(load_json(Path(args.site_json)))

    result = verify(master, candidates, groups, raw_counts, site_json_count)
    (out_dir / "verification.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print("[6/6] 검증 결과 저장: " + str(out_dir / "verification.json"))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["schema"]["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
