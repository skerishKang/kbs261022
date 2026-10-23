"""전수 카탈로그 통합 테스트.

수집 원본 → PR #16 마스터 → XLSX → 사이트 JSON 흐름을 실제 공개 샘플로 검증한다.
"""
from __future__ import annotations

import importlib.util
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[3]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from tools.full_catalog import build_full_catalog as BFC  # noqa: E402
from tools.full_catalog import build_site_data as BSD  # noqa: E402
from tools.product_master import dedupe as DD  # noqa: E402
from tools.product_master import schema as SC  # noqa: E402
from tools.product_master import validate as VA  # noqa: E402


def _load_pr17():
    path = REPO_ROOT / "tools" / "build_catalog_data.py"
    spec = importlib.util.spec_from_file_location("pr17_builder_test", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def samasound_raw(**over):
    raw = {
        "source_site": "samasound",
        "source_site_label": "Tri-A / 삼아사운드",
        "source_product_id": "3875",
        "source_category_id": "23",
        "source_category_name": "무선 바디팩용 마이크",
        "source_category_path": "홈 > 제품분류별 보기 > 마이크 > 무선마이크",
        "source_brand": "Shure",
        "source_item_code": "P0000SJU000A",
        "product_code": "P0000SJU",
        "product_name": "WL185MB/C-TQG",
        "model": "",
        "manufacturer": "Shure",
        "canonical_url": "https://samasound.co.kr/product/wl185mbc-tqg/3875/",
        "source_list_url": "/product/wl185mbc-tqg/3875/category/42/display/1/",
        "summary": "단일지향성 로우 프로파일 카디오이드 라발리에 컨덴서 마이크",
        "country_of_origin": "아시아 중국",
        "specs": json.dumps(
            {
                "상품명": "WL185MB/C-TQG",
                "제조사": "Shure",
                "원산지": "아시아 중국",
                "판매가": "가격 문의",
            },
            ensure_ascii=False,
        ),
        "image_urls": "https://www.samasound.co.kr/web/product/small/a.jpg|/web/upload/x.png",
        "option_summary": json.dumps({"main": [], "addproduct_options": [], "addproducts": [], "add_option_data": []}),
        "option_count": 0,
        "published_price": "",
        "published_price_text": "가격 문의",
        "published_price_state": "inquiry_price",
        "published_price_unavailable": 1,
        "availability_unverified": "{}",
        "extraction_method": "http_get_public_html",
        "verification_status": "verified_against_source_page",
        "updated_at": "",
    }
    raw.update(over)
    return raw


def samamusic_raw(**over):
    raw = {
        "source_site": "samamusic",
        "source_site_label": "Tri-B / 삼아스토어",
        "source_product_id": "5557",
        "source_category_id": "671",
        "source_category_name": "SHURE",
        "source_category_path": "",
        "source_brand": "SHURE",
        "source_item_code": "P0000SJU000A",
        "product_code": "P0000SJU",
        "product_name": "SHURE WL185MB/C-TQG 무선용 단일 지향성 핀마이크 블랙",
        "model": "WL185MB/C-TQG",
        "manufacturer": "",
        "canonical_url": "https://samastore.co.kr/product/shure-wl185mbc-tqg/5557/",
        "source_list_url": "/product/detail.html?product_no=5557&cate_no=671&display_group=1",
        "summary": "무선용 단일 지향성 핀마이크",
        "country_of_origin": "아시아 중국",
        "specs": json.dumps(
            {
                "상품명": "SHURE WL185MB/C-TQG 무선용 단일 지향성 핀마이크 블랙",
                "브랜드": "SHURE",
                "모델명": "WL185MB/C-TQG",
                "원산지": "아시아 중국",
                "판매가": "279,000원",
            },
            ensure_ascii=False,
        ),
        "image_urls": "https://www.samamusic.co.kr/web/product/small/b.jpg",
        "option_summary": json.dumps(
            {
                "main": [
                    {
                        "option_name": "색상 선택",
                        "values": ["클리어 [품절]", "블랙"],
                        "item_codes": ["P0000SJU000A", "P0000SJU000B"],
                    }
                ],
                "addproduct_options": [],
                "addproducts": [],
                "add_option_data": [],
            },
            ensure_ascii=False,
        ),
        "option_count": 1,
        "published_price": 279000,
        "published_price_text": "279,000원",
        "published_price_state": "numeric_price",
        "published_price_unavailable": 0,
        "availability_unverified": "{}",
        "extraction_method": "http_get_public_html",
        "verification_status": "verified_against_source_page",
        "updated_at": "",
    }
    raw.update(over)
    return raw


class TestToMasterRecord(unittest.TestCase):
    def test_samasound_brand_comes_from_category(self) -> None:
        record = BFC.to_master_record(samasound_raw(), "2026-10-22T00:00:00Z")
        assert record is not None
        # 브랜드 스펙행이 없으므로 brand_source 는 'category'
        self.assertEqual(record["brand"], "SHURE")
        self.assertEqual(record["brand_source"], "category")
        self.assertEqual(record["manufacturer"], "Shure")
        self.assertEqual(record["manufacturer_source"], "spec")
        self.assertIsNone(record["model"])

    def test_samamusic_brand_comes_from_spec(self) -> None:
        record = BFC.to_master_record(samamusic_raw(), "2026-10-22T00:00:00Z")
        assert record is not None
        self.assertEqual(record["brand"], "SHURE")
        self.assertEqual(record["brand_source"], "spec")
        self.assertIsNone(record["manufacturer"])
        self.assertEqual(record["manufacturer_source"], "none")
        self.assertEqual(record["model"], "WL185MB/C-TQG")

    def test_inquiry_price_has_no_number(self) -> None:
        record = BFC.to_master_record(samasound_raw(), "2026-10-22T00:00:00Z")
        assert record is not None
        self.assertIsNone(record["public_price"])
        self.assertEqual(record["price_type"], "inquiry")
        self.assertEqual(record["public_price_text"], "가격 문의")
        self.assertIsNone(record["dealer_price"])
        self.assertIsNone(record["availability"])

    def test_numeric_price_is_kept(self) -> None:
        record = BFC.to_master_record(samamusic_raw(), "2026-10-22T00:00:00Z")
        assert record is not None
        self.assertEqual(record["public_price"], 279000)
        self.assertEqual(record["price_type"], "public_selling_price")

    def test_sold_out_price_is_unknown_not_inquiry(self) -> None:
        record = BFC.to_master_record(
            samamusic_raw(published_price="", published_price_text="판매 완료", published_price_state="sold_out_sale_completed"),
            "2026-10-22T00:00:00Z",
        )
        assert record is not None
        self.assertIsNone(record["public_price"])
        self.assertEqual(record["price_type"], "unknown")
        self.assertIn("sold_out", record["quality_flags"])

    def test_relative_image_is_absolutised(self) -> None:
        record = BFC.to_master_record(samasound_raw(), "2026-10-22T00:00:00Z")
        assert record is not None
        self.assertIn("https://www.samasound.co.kr/web/upload/x.png", record["image_urls"])

    def test_options_are_parsed_with_soldout(self) -> None:
        record = BFC.to_master_record(samamusic_raw(), "2026-10-22T00:00:00Z")
        assert record is not None
        self.assertEqual(len(record["options"]), 1)
        values = record["options"][0]["values"]
        self.assertTrue(values[0]["soldout"])
        self.assertFalse(values[1]["soldout"])
        self.assertEqual(values[1]["item_code"], "P0000SJU000B")

    def test_category_namespace_is_site_scoped(self) -> None:
        a = BFC.to_master_record(samasound_raw(), "2026-10-22T00:00:00Z")
        b = BFC.to_master_record(samamusic_raw(), "2026-10-22T00:00:00Z")
        assert a is not None and b is not None
        self.assertEqual(a["category_namespace"], "samasound.cate_no")
        self.assertEqual(b["category_namespace"], "samamusic.cate_no")
        self.assertNotEqual(a["category_namespace"], b["category_namespace"])

    def test_record_passes_policy_validation(self) -> None:
        for raw in (samasound_raw(), samamusic_raw()):
            record = BFC.to_master_record(raw, "2026-10-22T00:00:00Z")
            assert record is not None
            self.assertEqual(VA.validate_record(record), [])
            self.assertEqual(VA.validate_policy(record), [])


class TestDedupeIntegration(unittest.TestCase):
    def test_cross_site_pair_is_confirmed(self) -> None:
        records = [
            BFC.to_master_record(samasound_raw(), "2026-10-22T00:00:00Z"),
            BFC.to_master_record(samamusic_raw(), "2026-10-22T00:00:00Z"),
        ]
        records = [r for r in records if r]
        candidates = DD.find_candidates(records)
        confirmed = [c for c in candidates if c.verdict == "confirmed"]
        self.assertEqual(len(confirmed), 1)
        groups = DD.build_groups(records, candidates)
        self.assertEqual(len(groups), 1)


class TestSheetBuilders(unittest.TestCase):
    def _records(self):
        return [
            r
            for r in (
                BFC.to_master_record(samasound_raw(), "2026-10-22T00:00:00Z"),
                BFC.to_master_record(samamusic_raw(), "2026-10-22T00:00:00Z"),
            )
            if r
        ]

    def test_unified_sheet_has_source_ids_and_blank_unknowns(self) -> None:
        rows = BFC.build_unified_sheet(self._records())
        self.assertEqual(len(rows), 2)
        for row in rows:
            self.assertIn("공급처원본ID", row)
            self.assertIn(row["공급처원본ID"], {"3875", "5557"})
            self.assertEqual(row["도매가"], "")
            self.assertEqual(row["재고"], "")
            self.assertEqual(row["매핑준비"], "아니오")
            self.assertEqual(row["네이버상품ID"], "")
            self.assertEqual(row["ERP품목코드"], "")

    def test_naver_sheet_marks_unpriced_items(self) -> None:
        rows = BFC.build_naver_sheet(self._records())
        by_id = {r["공급처원본ID"]: r for r in rows}
        self.assertEqual(by_id["samasound:3875"]["가격확인필요"], "예")
        self.assertEqual(by_id["samasound:3875"]["등록가능여부"], "아니오")
        self.assertEqual(by_id["samamusic:5557"]["가격확인필요"], "아니오")
        self.assertEqual(by_id["samamusic:5557"]["등록가능여부"], "예")
        self.assertEqual(by_id["samamusic:5557"]["네이버상품ID"], "")

    def test_erp_sheet_has_no_erp_codes(self) -> None:
        rows = BFC.build_erp_sheet(self._records())
        for row in rows:
            self.assertEqual(row["이카운트품목코드"], "")
            self.assertEqual(row["ERP매핑상태"], "미매핑")

    def test_brand_category_sheet_keeps_namespace(self) -> None:
        rows = BFC.build_brand_category_sheet(self._records())
        self.assertEqual(len(rows), 2)
        namespaces = {r["cate_no_namespace"] for r in rows}
        self.assertEqual(namespaces, {"samasound.cate_no", "samamusic.cate_no"})

    def test_supplier_sheet_preserves_original_ids(self) -> None:
        rows = BFC.build_supplier_sheet(
            self._records(),
            {
                "samasound:3875": samasound_raw(),
                "samamusic:5557": samamusic_raw(),
            },
        )
        self.assertEqual(len(rows), 2)
        by_id = {r["공급처원본ID"]: r for r in rows}
        self.assertEqual(by_id["3875"]["품목코드"], "P0000SJU000A")
        self.assertEqual(by_id["5557"]["원본브랜드"], "SHURE")
        self.assertEqual(by_id["3875"]["원본제조사"], "Shure")

    def test_supplier_sheet_collect_time_falls_back_to_extraction_time(self) -> None:
        """상세 덤프에 `updated_at` 이 없으면(실측 0건) 통합상품과 같은 수집 시각을 쓴다.

        커밋본 `공급처원본` 시트가 이 열을 빈 칸으로 남겼던 원인이 이 폴백 부재였다.
        폴백 값은 상품별 갱신 시각이 아니라 DB 생성(빌드) 시각이다.
        """
        records = self._records()
        rows = BFC.build_supplier_sheet(
            records,
            {"samasound:3875": samasound_raw(updated_at=""), "samamusic:5557": samamusic_raw(updated_at="")},
        )
        unified = {r["공급처원본ID"]: r for r in BFC.build_unified_sheet(records)}
        for row in rows:
            self.assertEqual(row["수집시각"], "2026-10-22T00:00:00Z")
            self.assertEqual(row["수집시각"], unified[row["공급처원본ID"]]["수집시각"])

    def test_supplier_sheet_prefers_source_updated_at(self) -> None:
        rows = BFC.build_supplier_sheet(
            self._records(),
            {"samasound:3875": samasound_raw(updated_at="2026-09-01T00:00:00Z"),
             "samamusic:5557": samamusic_raw()},
        )
        by_id = {r["공급처원본ID"]: r for r in rows}
        self.assertEqual(by_id["3875"]["수집시각"], "2026-09-01T00:00:00Z")
        self.assertEqual(by_id["5557"]["수집시각"], "2026-10-22T00:00:00Z")


class TestXlsx(unittest.TestCase):
    def test_workbook_opens_with_required_sheets(self) -> None:
        try:
            from openpyxl import load_workbook
        except ImportError:  # pragma: no cover
            self.skipTest("openpyxl 미설치")
        records = [
            r
            for r in (
                BFC.to_master_record(samasound_raw(), "2026-10-22T00:00:00Z"),
                BFC.to_master_record(samamusic_raw(), "2026-10-22T00:00:00Z"),
            )
            if r
        ]
        sheets = [
            ("통합상품", BFC.build_unified_sheet(records)),
            ("공급처원본", BFC.build_supplier_sheet(records, {})),
            ("브랜드_카테고리", BFC.build_brand_category_sheet(records)),
            ("중복후보", []),
            ("누락_검증현황", BFC.build_gap_sheet(records)),
            ("네이버스토어준비", BFC.build_naver_sheet(records)),
            ("이카운트품목매핑", BFC.build_erp_sheet(records)),
            ("옵션SKU", BFC.build_option_sheet(records)),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "db.xlsx"
            BFC.write_xlsx(path, sheets)
            self.assertTrue(path.exists())
            wb = load_workbook(path, read_only=True)
            self.assertEqual(
                set(wb.sheetnames),
                {
                    "통합상품",
                    "공급처원본",
                    "브랜드_카테고리",
                    "중복후보",
                    "누락_검증현황",
                    "네이버스토어준비",
                    "이카운트품목매핑",
                    "옵션SKU",
                },
            )
            ws = wb["통합상품"]
            rows = list(ws.iter_rows(values_only=True))
            self.assertEqual(len(rows), 3)  # 헤더 + 2
            wb.close()


class TestSiteDataMatch(unittest.TestCase):
    def test_verify_match_detects_mismatch(self) -> None:
        products = [{"id": "samasound.co.kr:1"}, {"id": "samamusic.co.kr:2"}]
        master = [
            {"source_site": "samasound", "source_product_id": "1"},
            {"source_site": "samamusic", "source_product_id": "3"},
        ]
        result = BSD.verify_match(products, master)
        self.assertTrue(result["count_match"])
        self.assertFalse(result["id_match"])
        self.assertEqual(result["missing_in_site_count"], 1)

    def test_verify_match_accepts_equal_sets(self) -> None:
        products = [{"id": "samasound.co.kr:1"}, {"id": "samamusic.co.kr:2"}]
        master = [
            {"source_site": "samasound", "source_product_id": "1"},
            {"source_site": "samamusic", "source_product_id": "2"},
        ]
        result = BSD.verify_match(products, master)
        self.assertTrue(result["count_match"])
        self.assertTrue(result["id_match"])
        self.assertEqual(result["duplicate_ids_in_site"], 0)


if __name__ == "__main__":
    unittest.main()
