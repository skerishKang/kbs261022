#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""CGI LOCAL 3 - Nara Market P1 original-notice verification exporter.

Reads the curated verification dataset (p1_verified_notices.json), re-derives
the review state at a cutoff timestamp, compares the requested item profile
against the read-only CGI catalog, and writes review CSV/XLSX, quote-review
items, an evidence manifest, a summary JSON and a Markdown brief.

The tool never submits a bid and never modifies collected notice data, the CGI
catalog, or LOCAL1/LOCAL2 artifacts.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import sys
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:  # pragma: no cover
    openpyxl = None

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DEFAULT_DATASET = os.path.join(REPO_ROOT, "tools", "procurement", "p1_verified_notices.json")
DEFAULT_CATALOG = os.path.join(REPO_ROOT, "data", "full_catalog", "product_db_compact.csv")
DEFAULT_OUT_DIR = os.path.join(REPO_ROOT, "reports", "data", "procurement")
DEFAULT_AS_OF = "2026-10-23T12:10:00+09:00"

KST = timezone(timedelta(hours=9))

PROFILES = {
    "broadcast_install": {
        "label": "방송·음향 설치",
        "categories": [
            "스피커", "설치용 스피커", "PA 스피커", "액티브 스피커", "라우드 스피커",
            "스테이지 모니터 스피커", "서브우퍼", "디지털 믹서", "마이크",
            "유선 마이크", "무선 마이크", "파워앰프",
        ],
    },
    "speaker_system": {
        "label": "스피커 시스템",
        "categories": [
            "스피커", "설치용 스피커", "PA 스피커", "액티브 스피커", "라우드 스피커",
            "스테이지 모니터 스피커", "서브우퍼",
        ],
    },
    "mixer": {
        "label": "디지털 믹서",
        "categories": ["디지털 믹서"],
    },
    "projector": {
        "label": "프로젝터",
        "categories": [],
    },
}

SALES_REVIEW = {
    "ACTIVE": "검토 가능(진행)",
    "EXPIRED": "검토 종료(마감)",
    "NOT_ELIGIBLE": "제외(비경쟁)",
    "UNVERIFIED": "보류(원문 미확인)",
}

STATE_ORDER = {"ACTIVE": 0, "EXPIRED": 1, "NOT_ELIGIBLE": 2, "UNVERIFIED": 3}

MAIN_COLUMNS = [
    "p1_priority", "bid_no", "verification_status", "review_state", "sales_review",
    "eligibility", "source_system", "source_url", "portal_path", "checked_at",
    "title", "notice_type", "notice_no", "ref_no", "posted_at",
    "agency", "demand_agency", "bid_method", "award_method", "contract_method",
    "joint_supply", "region_restriction", "region_detail",
    "license_restriction", "license_detail", "product_registration",
    "estimated_price", "price_basis", "deadline_kind", "deadline_label",
    "deadline", "cancellation_status", "opening", "delivery", "warranty",
    "qualification_verified", "qualifications_count",
    "documents_verified", "documents_checked_at",
    "documents_count", "documents_extracted_count",
    "cgi_fit_level", "cgi_category_profile", "cgi_matched_products",
    "quote_review_count", "next_action", "status_reason",
]

QUOTE_COLUMNS = ["bid_no", "title", "review_state", "sales_review", "seq", "item"]
QUALIFICATION_COLUMNS = ["bid_no", "seq", "text", "source"]
DOCUMENT_COLUMNS = [
    "bid_no", "kind", "name", "size_bytes", "sha256", "text_extracted",
    "extract_chars", "verified_topics", "extract_files", "note",
]
EVIDENCE_COLUMNS = ["bid_no", "topic", "quote", "source"]
MATCH_COLUMNS = [
    "bid_no", "title", "cgi_fit_level", "category_profile", "profile_label",
    "matched_product_count", "matched_categories", "sales_status_counts",
    "sample_products", "note",
]


class DatasetError(ValueError):
    """Raised when the curated dataset is inconsistent or unsupported."""


def parse_dt(value):
    if value in (None, ""):
        return None
    text = str(value).strip()
    if text.endswith("Z"):
        text = text[:-1] + "+00:00"
    try:
        dt = datetime.fromisoformat(text)
    except ValueError as exc:
        raise DatasetError(f"invalid timestamp: {value!r}") from exc
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=KST)
    return dt


def as_kst(dt):
    return dt.astimezone(KST)


def compute_review_state(notice, as_of):
    """Derive ACTIVE/EXPIRED/NOT_ELIGIBLE/UNVERIFIED from verified facts."""
    if notice.get("verification_status") != "VERIFIED":
        return "UNVERIFIED"
    eligibility = notice.get("eligibility")
    if eligibility == "NOT_ELIGIBLE":
        return "NOT_ELIGIBLE"
    deadline = parse_dt(notice.get("deadline"))
    if deadline is None:
        return "UNVERIFIED"
    if eligibility == "OPEN_COMPETITIVE":
        return "ACTIVE" if deadline > as_of else "EXPIRED"
    if eligibility == "CLOSED":
        if deadline > as_of:
            raise DatasetError(
                f"{notice.get('bid_no')}: eligibility=CLOSED but deadline "
                f"{deadline.isoformat()} is after as_of {as_of.isoformat()}"
            )
        return "EXPIRED"
    if eligibility == "UNKNOWN":
        return "UNVERIFIED"
    raise DatasetError(f"{notice.get('bid_no')}: unknown eligibility {eligibility!r}")


def _check_sha256(value, label):
    if not value or len(value) != 64 or any(c not in "0123456789abcdef" for c in value.lower()):
        raise DatasetError(f"{label}: invalid sha256")
    return value.lower()


def validate_dataset(dataset, as_of):
    notices = dataset.get("notices")
    if not isinstance(notices, list) or not notices:
        raise DatasetError("dataset has no notices")
    seen = set()
    errors = []
    for idx, notice in enumerate(notices, start=1):
        bid_no = notice.get("bid_no")
        if not bid_no:
            errors.append(f"record #{idx}: missing bid_no")
            continue
        if bid_no in seen:
            errors.append(f"{bid_no}: duplicate bid_no")
        seen.add(bid_no)

        status = notice.get("verification_status")
        if status not in ("VERIFIED", "UNVERIFIED"):
            errors.append(f"{bid_no}: invalid verification_status {status!r}")
            continue
        declared = notice.get("review_state")
        if declared not in SALES_REVIEW:
            errors.append(f"{bid_no}: invalid review_state {declared!r}")
        computed = compute_review_state(notice, as_of)
        if declared != computed:
            errors.append(
                f"{bid_no}: declared review_state {declared} != computed {computed} "
                f"at as_of {as_of.isoformat()}"
            )

        if status == "VERIFIED":
            url = notice.get("source_url") or ""
            if not url.startswith("https://"):
                errors.append(f"{bid_no}: verified record needs an https source_url")
            checked_at = parse_dt(notice.get("checked_at"))
            if checked_at is None or checked_at.utcoffset() != timedelta(hours=9):
                errors.append(f"{bid_no}: checked_at must be a KST (+09:00) timestamp")
            deadline = parse_dt(notice.get("deadline"))
            if deadline is None:
                errors.append(f"{bid_no}: verified record has no deadline")
            elif deadline.utcoffset() != timedelta(hours=9):
                errors.append(f"{bid_no}: deadline must be a KST (+09:00) timestamp")
            if not notice.get("cancellation_status"):
                errors.append(f"{bid_no}: verified record has no cancellation_status")
        else:
            if notice.get("deadline") is not None:
                errors.append(f"{bid_no}: unverified record must not assert a deadline")
            if notice.get("qualification_verified"):
                errors.append(f"{bid_no}: unverified record must not assert qualifications")
            for evidence in notice.get("evidence_quotes") or []:
                if evidence.get("topic") != "원문 접근":
                    errors.append(
                        f"{bid_no}: unverified record may only quote access blockers, "
                        f"found topic {evidence.get('topic')!r}"
                    )
            if not notice.get("status_reason"):
                errors.append(f"{bid_no}: unverified record needs status_reason")

        if notice.get("qualification_verified") and not notice.get("qualifications"):
            errors.append(f"{bid_no}: qualification_verified without qualification rows")
        if notice.get("documents_verified"):
            docs = notice.get("documents") or []
            if not any(doc.get("text_extracted") for doc in docs):
                errors.append(f"{bid_no}: documents_verified without an extracted document")

        for doc in notice.get("documents") or []:
            if not doc.get("text_extracted"):
                continue
            if not doc.get("extract_chars"):
                errors.append(f"{bid_no}: extracted document has no extract_chars")
            files = doc.get("extract_files") or []
            if not files:
                errors.append(f"{bid_no}: extracted document has no committed extract")
            for entry in files:
                rel = entry.get("file")
                if not rel:
                    errors.append(f"{bid_no}: extract entry missing file")
                    continue
                path = os.path.join(REPO_ROOT, rel)
                if not os.path.exists(path):
                    errors.append(f"{bid_no}: missing extract file {rel}")
                    continue
                with open(path, "rb") as f:
                    payload = f.read()
                if entry.get("size_bytes") != len(payload):
                    errors.append(f"{bid_no}: size mismatch for {rel}")
                _check_sha256(entry.get("sha256"), f"{bid_no}: {rel}")
                if hashlib.sha256(payload).hexdigest() != entry["sha256"].lower():
                    errors.append(f"{bid_no}: sha256 mismatch for {rel}")

    if len(seen) != len(notices):
        errors.append("bid_no values are not unique")
    if errors:
        raise DatasetError("dataset validation failed:\n- " + "\n- ".join(errors))
    return notices


def load_dataset(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def load_catalog(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    if not rows:
        raise DatasetError("CGI catalog is empty")
    return rows


def build_catalog_index(rows):
    by_cat = defaultdict(list)
    for row in rows:
        by_cat[row.get("category_name") or ""].append(row)
    return by_cat


def _status_priority(row):
    order = {"판매중": 0, "재고미관리": 1, "미확인": 2, "판매완료": 3, "품절(판매중지)": 4}
    return order.get(row.get("sales_status") or "", 9)


def match_notice(notice, by_cat):
    fit = notice.get("cgi_fit") or {}
    profile = fit.get("category_profile")
    base = {
        "bid_no": notice["bid_no"],
        "title": notice.get("title") or "",
        "cgi_fit_level": fit.get("level") or "",
        "category_profile": profile or "",
        "profile_label": "",
        "matched_product_count": "",
        "matched_categories": "",
        "sales_status_counts": "",
        "sample_products": "",
        "note": fit.get("note") or "",
    }
    if not profile:
        return base
    if profile not in PROFILES:
        raise DatasetError(f"{notice['bid_no']}: unknown category profile {profile!r}")
    spec = PROFILES[profile]
    base["profile_label"] = spec["label"]

    products = []
    breakdown = []
    for cat in spec["categories"]:
        items = by_cat.get(cat, [])
        if items:
            breakdown.append(f"{cat}:{len(items)}")
        products.extend(items)

    base["matched_product_count"] = len(products)
    base["matched_categories"] = "; ".join(breakdown)
    status_counts = Counter(row.get("sales_status") or "미확인" for row in products)
    base["sales_status_counts"] = "; ".join(
        f"{key}:{value}" for key, value in status_counts.most_common()
    )
    examples = sorted(products, key=_status_priority)[:3]
    base["sample_products"] = "; ".join(
        "{}{}{}".format(
            row.get("product_name") or "",
            f" / {row.get('brand')}" if row.get("brand") else "",
            f" / {row.get('public_price_text')}" if row.get("public_price_text") else "",
        )
        for row in examples
    )
    return base


def build_rows(notices, by_cat, as_of):
    main_rows = []
    quote_rows = []
    qualification_rows = []
    document_rows = []
    evidence_rows = []
    match_rows = []
    for priority, notice in enumerate(notices, start=1):
        state = compute_review_state(notice, as_of)
        fit = notice.get("cgi_fit") or {}
        match = match_notice(notice, by_cat)
        match_rows.append(match)

        documents = notice.get("documents") or []
        qualifications = notice.get("qualifications") or []
        quotes = notice.get("quote_review_items") or []
        main_rows.append({
            "p1_priority": priority,
            "bid_no": notice["bid_no"],
            "verification_status": notice.get("verification_status") or "",
            "review_state": state,
            "sales_review": SALES_REVIEW[state],
            "eligibility": notice.get("eligibility") or "",
            "source_system": notice.get("source_system") or "",
            "source_url": notice.get("source_url") or "",
            "portal_path": notice.get("portal_path") or "",
            "checked_at": notice.get("checked_at") or "",
            "title": notice.get("title") or "",
            "notice_type": notice.get("notice_type") or "",
            "notice_no": notice.get("notice_no") or "",
            "ref_no": notice.get("ref_no") or "",
            "posted_at": notice.get("posted_at") or "",
            "agency": notice.get("agency") or "",
            "demand_agency": notice.get("demand_agency") or "",
            "bid_method": notice.get("bid_method") or "",
            "award_method": notice.get("award_method") or "",
            "contract_method": notice.get("contract_method") or "",
            "joint_supply": notice.get("joint_supply") or "",
            "region_restriction": notice.get("region_restriction") or "",
            "region_detail": notice.get("region_detail") or "",
            "license_restriction": notice.get("license_restriction") or "",
            "license_detail": notice.get("license_detail") or "",
            "product_registration": notice.get("product_registration") or "",
            "estimated_price": notice.get("estimated_price") if notice.get("estimated_price") is not None else "",
            "price_basis": notice.get("price_basis") or "",
            "deadline_kind": notice.get("deadline_kind") or "",
            "deadline_label": notice.get("deadline_label") or "",
            "deadline": notice.get("deadline") or "",
            "cancellation_status": notice.get("cancellation_status") or "",
            "opening": notice.get("opening") or "",
            "delivery": notice.get("delivery") or "",
            "warranty": notice.get("warranty") or "",
            "qualification_verified": bool(notice.get("qualification_verified")),
            "qualifications_count": len(qualifications),
            "documents_verified": bool(notice.get("documents_verified")),
            "documents_checked_at": notice.get("documents_checked_at") or "",
            "documents_count": len(documents),
            "documents_extracted_count": sum(1 for doc in documents if doc.get("text_extracted")),
            "cgi_fit_level": fit.get("level") or "",
            "cgi_category_profile": fit.get("category_profile") or "",
            "cgi_matched_products": match["matched_product_count"],
            "quote_review_count": len(quotes),
            "next_action": notice.get("next_action") or "",
            "status_reason": notice.get("status_reason") or "",
        })

        for seq, item in enumerate(quotes, start=1):
            quote_rows.append({
                "bid_no": notice["bid_no"],
                "title": notice.get("title") or "",
                "review_state": state,
                "sales_review": SALES_REVIEW[state],
                "seq": seq,
                "item": item,
            })
        for seq, qual in enumerate(qualifications, start=1):
            qualification_rows.append({
                "bid_no": notice["bid_no"],
                "seq": seq,
                "text": qual.get("text") or "",
                "source": qual.get("source") or "",
            })
        for doc in documents:
            document_rows.append({
                "bid_no": notice["bid_no"],
                "kind": doc.get("kind") or "",
                "name": doc.get("name") or "",
                "size_bytes": doc.get("size_bytes") if doc.get("size_bytes") is not None else "",
                "sha256": doc.get("sha256") or "",
                "text_extracted": bool(doc.get("text_extracted")),
                "extract_chars": doc.get("extract_chars") or 0,
                "verified_topics": "; ".join(doc.get("verified_topics") or []),
                "extract_files": "; ".join(
                    entry.get("file") or "" for entry in doc.get("extract_files") or []
                ),
                "note": doc.get("note") or "",
            })
        for evidence in notice.get("evidence_quotes") or []:
            evidence_rows.append({
                "bid_no": notice["bid_no"],
                "topic": evidence.get("topic") or "",
                "quote": evidence.get("quote") or "",
                "source": evidence.get("source") or "",
            })
    return {
        "main": main_rows,
        "quotes": quote_rows,
        "qualifications": qualification_rows,
        "documents": document_rows,
        "evidence": evidence_rows,
        "matches": match_rows,
    }


def build_summary(dataset, notices, rows, by_cat, catalog_rows, as_of):
    by_state = Counter(row["review_state"] for row in rows["main"])
    by_status = Counter(row["verification_status"] for row in rows["main"])
    verified = [row for row in rows["main"] if row["verification_status"] == "VERIFIED"]
    qual_verified = [row for row in rows["main"] if row["qualification_verified"]]
    docs_verified = [row for row in rows["main"] if row["documents_verified"]]

    def brief(row):
        return {
            "bid_no": row["bid_no"],
            "title": row["title"],
            "agency": row["agency"],
            "deadline": row["deadline"],
            "deadline_label": row["deadline_label"],
            "estimated_price": row["estimated_price"],
            "source_url": row["source_url"],
            "next_action": row["next_action"],
        }

    active = sorted(
        (row for row in rows["main"] if row["review_state"] == "ACTIVE"),
        key=lambda row: row["deadline"],
    )
    expired = sorted(
        (row for row in rows["main"] if row["review_state"] == "EXPIRED"),
        key=lambda row: row["deadline"],
        reverse=True,
    )
    not_eligible = [row for row in rows["main"] if row["review_state"] == "NOT_ELIGIBLE"]
    unverified = [row for row in rows["main"] if row["review_state"] == "UNVERIFIED"]

    blockers = []
    for notice in notices:
        state = compute_review_state(notice, as_of)
        if state == "UNVERIFIED":
            blockers.append({
                "bid_no": notice["bid_no"],
                "kind": "PORTAL_OR_SEARCH",
                "review_state": state,
                "detail": notice.get("status_reason") or "",
            })
        elif state == "ACTIVE" and not notice.get("documents_verified"):
            reasons = [
                doc.get("note") for doc in notice.get("documents") or [] if doc.get("note")
            ]
            blockers.append({
                "bid_no": notice["bid_no"],
                "kind": "DOCUMENT",
                "review_state": state,
                "detail": "; ".join(reasons) or "첨부 원문 미확인",
            })

    profile_counts = {}
    for name, spec in PROFILES.items():
        products = []
        for cat in spec["categories"]:
            products.extend(by_cat.get(cat, []))
        profile_counts[name] = {
            "label": spec["label"],
            "categories": spec["categories"],
            "catalog_products": len(products),
            "sales_status": dict(Counter(row.get("sales_status") or "미확인" for row in products)),
        }

    return {
        "dataset": dataset.get("dataset"),
        "issue": dataset.get("issue"),
        "as_of": as_of.isoformat(),
        "total_notices": len(rows["main"]),
        "by_review_state": {key: by_state.get(key, 0) for key in STATE_ORDER},
        "by_verification_status": {
            "VERIFIED": by_status.get("VERIFIED", 0),
            "UNVERIFIED": by_status.get("UNVERIFIED", 0),
        },
        "verified_detail": len(verified),
        "verified_deadline": sum(1 for row in verified if row["deadline"]),
        "qualification_verified": len(qual_verified),
        "documents_verified": len(docs_verified),
        "quote_review_items_total": len(rows["quotes"]),
        "active": [brief(row) for row in active],
        "active_count": len(active),
        "expired": [brief(row) for row in expired],
        "expired_count": len(expired),
        "not_eligible": [brief(row) for row in not_eligible],
        "not_eligible_count": len(not_eligible),
        "unverified": [brief(row) for row in unverified],
        "unverified_count": len(unverified),
        "access_blockers": blockers,
        "cgi_catalog": {
            "source": "data/full_catalog/product_db_compact.csv",
            "products_total": len(catalog_rows),
            "caveat": (
                "카테고리 단위 참고 비교다. 카탈로그 전 제품의 cgi_sale_status가 unapproved로 "
                "기록되어 있어 판매승인 필터로 사용할 수 없고, sales_status도 재고미관리·품절이 "
                "대부분이므로 실제 견적 가능 모델은 별도 공급 확인이 필요하다."
            ),
            "profiles": profile_counts,
        },
        "policy": (
            "실제 원문(공고 상세 또는 첨부)에서 확인된 공고만 영업 검토 대상으로 확정한다. "
            "원문을 확인하지 못한 공고는 UNVERIFIED로 보류하며 마감일·자격을 추정하지 않는다."
        ),
    }


def write_csv(path, rows, columns):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def _style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for column in ws.columns:
        length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        letter = column[0].column_letter
        ws.column_dimensions[letter].width = min(max(12, length + 2), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_xlsx(path, rows, summary):
    if openpyxl is None:
        raise DatasetError("openpyxl is required to write the XLSX output")
    wb = openpyxl.Workbook()
    sheets = [
        ("P1_Verification", MAIN_COLUMNS, rows["main"]),
        ("Quote_Review", QUOTE_COLUMNS, rows["quotes"]),
        ("Qualifications", QUALIFICATION_COLUMNS, rows["qualifications"]),
        ("Documents", DOCUMENT_COLUMNS, rows["documents"]),
        ("Evidence_Quotes", EVIDENCE_COLUMNS, rows["evidence"]),
        ("CGI_Matches", MATCH_COLUMNS, rows["matches"]),
    ]
    first = True
    for title, columns, data_rows in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = title
        ws.append(columns)
        for row in data_rows:
            ws.append([row.get(column, "") for column in columns])
        _style_sheet(ws)

    ws = wb.create_sheet("Summary")
    ws.append(["metric", "value"])
    ws.append(["as_of", summary["as_of"]])
    ws.append(["total_notices", summary["total_notices"]])
    for key in STATE_ORDER:
        ws.append([f"review_state:{key}", summary["by_review_state"][key]])
    ws.append(["verified_detail", summary["verified_detail"]])
    ws.append(["verified_deadline", summary["verified_deadline"]])
    ws.append(["qualification_verified", summary["qualification_verified"]])
    ws.append(["documents_verified", summary["documents_verified"]])
    ws.append(["quote_review_items_total", summary["quote_review_items_total"]])
    ws.append(["catalog_products_total", summary["cgi_catalog"]["products_total"]])
    _style_sheet(ws)
    wb.save(path)


def write_manifest(path, dataset, notices, as_of):
    payload = {
        "dataset": dataset.get("dataset"),
        "issue": dataset.get("issue"),
        "as_of": as_of.isoformat(),
        "source_dataset": "tools/procurement/p1_verified_notices.json",
        "notice_count": len(notices),
        "notices": [],
    }
    for notice in notices:
        entry = {
            "bid_no": notice["bid_no"],
            "verification_status": notice["verification_status"],
            "review_state": compute_review_state(notice, as_of),
            "source_system": notice.get("source_system"),
            "source_url": notice.get("source_url"),
            "portal_path": notice.get("portal_path"),
            "checked_at": notice.get("checked_at"),
            "deadline": notice.get("deadline"),
            "deadline_label": notice.get("deadline_label"),
            "cancellation_status": notice.get("cancellation_status"),
            "qualification_verified": bool(notice.get("qualification_verified")),
            "documents_verified": bool(notice.get("documents_verified")),
            "documents_checked_at": notice.get("documents_checked_at"),
            "evidence_quotes": notice.get("evidence_quotes") or [],
            "documents": [
                {
                    "name": doc.get("name"),
                    "kind": doc.get("kind"),
                    "size_bytes": doc.get("size_bytes"),
                    "sha256": doc.get("sha256"),
                    "text_extracted": bool(doc.get("text_extracted")),
                    "extract_chars": doc.get("extract_chars") or 0,
                    "extract_files": doc.get("extract_files") or [],
                    "note": doc.get("note"),
                }
                for doc in notice.get("documents") or []
            ],
        }
        payload["notices"].append(entry)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


def _md_cell(value):
    if value in (None, ""):
        return "-"
    return str(value).replace("|", "\\|").replace("\n", " ")


def _md_table(columns, rows):
    lines = ["| " + " | ".join(columns) + " |", "|" + "|".join(["---"] * len(columns)) + "|"]
    for row in rows:
        lines.append("| " + " | ".join(_md_cell(row.get(column)) for column in columns) + " |")
    return "\n".join(lines)


def _won(value):
    if value in (None, ""):
        return "-"
    return f"{int(value):,}"


def write_markdown(path, summary, rows):
    by_no = {row["bid_no"]: row for row in rows["main"]}
    matches = {row["bid_no"]: row for row in rows["matches"]}
    lines = [
        "# CGI LOCAL 3 — 나라장터 P1 14건 원문 검증",
        "",
        f"- 기준시점(as_of): `{summary['as_of']}`",
        "- 검증 데이터: `tools/procurement/p1_verified_notices.json`",
        "- 원칙: 실제 원문(공고 상세 또는 첨부)에서 확인된 공고만 영업 검토 대상으로 확정한다.",
        "- 실제 입찰서 제출·투찰은 수행하지 않았다.",
        "",
        "## 요약",
        "",
        "| 구분 | 건수 |",
        "|---|---|",
        f"| P1 전체 | {summary['total_notices']} |",
        f"| 진행 가능(ACTIVE) | {summary['by_review_state']['ACTIVE']} |",
        f"| 마감(EXPIRED) | {summary['by_review_state']['EXPIRED']} |",
        f"| 제외(비경쟁/NOT_ELIGIBLE) | {summary['by_review_state']['NOT_ELIGIBLE']} |",
        f"| 원문 미확인(UNVERIFIED) | {summary['by_review_state']['UNVERIFIED']} |",
        f"| 상세 검증 완료 | {summary['verified_detail']} |",
        f"| 마감일 검증 완료 | {summary['verified_deadline']} |",
        f"| 참가자격 검증 완료 | {summary['qualification_verified']} |",
        f"| 첨부문서 검증 완료 | {summary['documents_verified']} |",
        "",
        "## 진행 가능(ACTIVE)",
        "",
    ]
    if summary["active"]:
        active_rows = []
        for entry in summary["active"]:
            row = by_no[entry["bid_no"]]
            match = matches.get(row["bid_no"], {})
            active_rows.append({
                "마감(KST)": row["deadline"],
                "공고번호": row["bid_no"],
                "공고명": row["title"],
                "발주기관": row["agency"],
                "예정가격(원)": _won(row["estimated_price"]),
                "자격": "검증" if row["qualification_verified"] else "미검증",
                "첨부": "검증" if row["documents_verified"] else "미검증",
                "CGI 적합도": row["cgi_fit_level"],
                "CGI 참고품목": match.get("matched_product_count", ""),
                "다음 조치": row["next_action"],
            })
        lines.append(_md_table(list(active_rows[0].keys()), active_rows))
    else:
        lines.append("없음.")
    lines += ["", "## 마감(EXPIRED)", ""]
    if summary["expired"]:
        expired_rows = []
        for entry in summary["expired"]:
            row = by_no[entry["bid_no"]]
            expired_rows.append({
                "마감(KST)": row["deadline"],
                "공고번호": row["bid_no"],
                "공고명": row["title"],
                "발주기관": row["agency"],
                "예정가격(원)": _won(row["estimated_price"]),
                "비고": row["deadline_label"],
            })
        lines.append(_md_table(list(expired_rows[0].keys()), expired_rows))
    else:
        lines.append("없음.")
    lines += ["", "## 제외(비경쟁)", ""]
    if summary["not_eligible"]:
        rows_ = []
        for entry in summary["not_eligible"]:
            row = by_no[entry["bid_no"]]
            rows_.append({
                "공고번호": row["bid_no"],
                "공고명": row["title"],
                "사유": row["status_reason"],
                "다음 조치": row["next_action"],
            })
        lines.append(_md_table(list(rows_[0].keys()), rows_))
    else:
        lines.append("없음.")
    lines += ["", "## 원문 미확인(UNVERIFIED)", ""]
    if summary["unverified"]:
        rows_ = []
        for entry in summary["unverified"]:
            row = by_no[entry["bid_no"]]
            rows_.append({
                "공고번호": row["bid_no"],
                "기존 분류 제목": row["title"],
                "사유": row["status_reason"],
                "다음 조치": row["next_action"],
            })
        lines.append(_md_table(list(rows_[0].keys()), rows_))
    else:
        lines.append("없음.")
    lines += [
        "",
        "## CGI 카탈로그 비교",
        "",
        f"- 카탈로그: `{summary['cgi_catalog']['source']}` (읽기 전용, {summary['cgi_catalog']['products_total']:,}개)",
        f"- 주의: {summary['cgi_catalog']['caveat']}",
        "",
    ]
    profile_rows = []
    for name, profile in summary["cgi_catalog"]["profiles"].items():
        profile_rows.append({
            "프로필": name,
            "설명": profile["label"],
            "카테고리": ", ".join(profile["categories"]) or "없음",
            "참고 제품 수": profile["catalog_products"],
        })
    lines.append(_md_table(list(profile_rows[0].keys()), profile_rows))
    lines += ["", "## 공고별 CGI 비교", ""]
    match_rows = []
    for row in rows["main"]:
        match = matches.get(row["bid_no"], {})
        match_rows.append({
            "공고번호": row["bid_no"],
            "검토상태": row["review_state"],
            "CGI 적합도": row["cgi_fit_level"],
            "프로필": match.get("category_profile") or "-",
            "참고 제품 수": match.get("matched_product_count", ""),
            "비고": match.get("note", ""),
        })
    lines.append(_md_table(list(match_rows[0].keys()), match_rows))
    lines += ["", "## 근거 문서", ""]
    document_rows = []
    for row in rows["documents"]:
        if row["text_extracted"]:
            document_rows.append({
                "공고번호": row["bid_no"],
                "문서": row["name"],
                "추출문자": row["extract_chars"],
                "SHA256(앞 12자)": (row["sha256"] or "")[:12],
                "추출본": row["extract_files"],
            })
    if document_rows:
        lines.append(_md_table(list(document_rows[0].keys()), document_rows))
    else:
        lines.append("없음.")
    lines += [
        "",
        "## 접근 차단·미검증 사유",
        "",
    ]
    for blocker in summary["access_blockers"]:
        lines.append(f"- `{blocker['bid_no']}` ({blocker['kind']}): {blocker['detail']}")
    lines += [
        "",
        "## 산출물",
        "",
        "- `reports/data/procurement/p1_original_verification.csv` — 공고별 검토 요약",
        "- `reports/data/procurement/p1_original_verification.xlsx` — 검토 요약·견적검토·자격·문서·근거·CGI 시트",
        "- `reports/data/procurement/p1_quote_review_items.csv` — 공고별 견적 검토 항목",
        "- `reports/data/procurement/p1_cgi_product_matches.csv` — CGI 카탈로그 참고 비교",
        "- `reports/data/procurement/p1_evidence_manifest.json` — 공고별 근거 URL·해시",
        "- `reports/data/procurement/p1_summary.json` — 집계",
        "- `reports/data/procurement/p1_evidence/` — 첨부 원문 텍스트 추출본",
        "",
    ]
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


def generate(dataset_path, catalog_path, out_dir, as_of):
    as_of = as_kst(parse_dt(as_of))
    dataset = load_dataset(dataset_path)
    notices = validate_dataset(dataset, as_of)
    catalog_rows = load_catalog(catalog_path)
    by_cat = build_catalog_index(catalog_rows)
    rows = build_rows(notices, by_cat, as_of)
    summary = build_summary(dataset, notices, rows, by_cat, catalog_rows, as_of)

    os.makedirs(out_dir, exist_ok=True)
    write_csv(os.path.join(out_dir, "p1_original_verification.csv"), rows["main"], MAIN_COLUMNS)
    write_csv(os.path.join(out_dir, "p1_quote_review_items.csv"), rows["quotes"], QUOTE_COLUMNS)
    write_csv(os.path.join(out_dir, "p1_cgi_product_matches.csv"), rows["matches"], MATCH_COLUMNS)
    write_xlsx(os.path.join(out_dir, "p1_original_verification.xlsx"), rows, summary)
    write_manifest(
        os.path.join(out_dir, "p1_evidence_manifest.json"), dataset, notices, as_of
    )
    with open(os.path.join(out_dir, "p1_summary.json"), "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
    write_markdown(
        os.path.join(out_dir, "p1_original_verification.md"), summary, rows
    )
    return summary


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dataset", default=DEFAULT_DATASET)
    parser.add_argument("--catalog", default=DEFAULT_CATALOG)
    parser.add_argument("--out-dir", default=DEFAULT_OUT_DIR)
    parser.add_argument("--as-of", default=DEFAULT_AS_OF)
    args = parser.parse_args(argv)
    try:
        summary = generate(args.dataset, args.catalog, args.out_dir, args.as_of)
    except DatasetError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    printable = {
        "as_of": summary["as_of"],
        "by_review_state": summary["by_review_state"],
        "by_verification_status": summary["by_verification_status"],
        "verified_detail": summary["verified_detail"],
        "verified_deadline": summary["verified_deadline"],
        "qualification_verified": summary["qualification_verified"],
        "documents_verified": summary["documents_verified"],
        "quote_review_items_total": summary["quote_review_items_total"],
        "output_dir": args.out_dir,
    }
    print(json.dumps(printable, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
