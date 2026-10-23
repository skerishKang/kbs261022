import csv
import json
import os
import sys
import tempfile
import unittest
from datetime import datetime

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
sys.path.insert(0, os.path.join(REPO_ROOT, "tools", "procurement"))
import g2b_p1_verification as p1

DATASET = os.path.join(REPO_ROOT, "tools", "procurement", "p1_verified_notices.json")
CATALOG = os.path.join(REPO_ROOT, "data", "full_catalog", "product_db_compact.csv")
AS_OF = p1.DEFAULT_AS_OF


def sample_notice(**overrides):
    notice = {
        "bid_no": "RTEST000001",
        "verification_status": "VERIFIED",
        "review_state": "ACTIVE",
        "eligibility": "OPEN_COMPETITIVE",
        "source_url": "https://www.g2b.go.kr/example",
        "checked_at": "2026-10-23T10:00:00+09:00",
        "deadline": "2026-10-30T10:00:00+09:00",
        "cancellation_status": "취소 공고 아님",
        "qualification_verified": False,
        "documents_verified": False,
    }
    notice.update(overrides)
    return notice


class TestReviewState(unittest.TestCase):
    def test_active_before_deadline(self):
        as_of = p1.parse_dt("2026-10-23T12:00:00+09:00")
        notice = sample_notice(review_state=None)
        self.assertEqual(p1.compute_review_state(notice, as_of), "ACTIVE")

    def test_expired_after_deadline(self):
        as_of = p1.parse_dt("2026-11-01T12:00:00+09:00")
        notice = sample_notice(review_state=None)
        self.assertEqual(p1.compute_review_state(notice, as_of), "EXPIRED")

    def test_not_eligible_regardless_of_deadline(self):
        as_of = p1.parse_dt("2026-10-23T12:00:00+09:00")
        notice = sample_notice(
            review_state="NOT_ELIGIBLE", eligibility="NOT_ELIGIBLE"
        )
        self.assertEqual(p1.compute_review_state(notice, as_of), "NOT_ELIGIBLE")

    def test_unverified_has_no_state(self):
        as_of = p1.parse_dt("2026-10-23T12:00:00+09:00")
        notice = sample_notice(
            verification_status="UNVERIFIED", review_state="UNVERIFIED", deadline=None
        )
        self.assertEqual(p1.compute_review_state(notice, as_of), "UNVERIFIED")

    def test_closed_with_future_deadline_is_rejected(self):
        as_of = p1.parse_dt("2026-10-23T12:00:00+09:00")
        notice = sample_notice(eligibility="CLOSED")
        with self.assertRaises(p1.DatasetError):
            p1.compute_review_state(notice, as_of)


class TestDatasetFacts(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dataset = p1.load_dataset(DATASET)
        cls.as_of = p1.as_kst(p1.parse_dt(AS_OF))
        cls.notices = p1.validate_dataset(cls.dataset, cls.as_of)
        cls.by_no = {notice["bid_no"]: notice for notice in cls.notices}

    def test_fourteen_unique_notices(self):
        self.assertEqual(len(self.notices), 14)
        self.assertEqual(len(self.by_no), 14)

    def test_verified_and_deadline_counts(self):
        verified = [n for n in self.notices if n["verification_status"] == "VERIFIED"]
        self.assertEqual(len(verified), 11)
        self.assertTrue(all(n.get("deadline") for n in verified))

    def test_all_verified_deadlines_are_kst(self):
        for notice in self.notices:
            if notice["verification_status"] != "VERIFIED":
                continue
            deadline = p1.parse_dt(notice["deadline"])
            self.assertEqual(deadline.utcoffset().total_seconds(), 9 * 3600)
            checked = p1.parse_dt(notice["checked_at"])
            self.assertEqual(checked.utcoffset().total_seconds(), 9 * 3600)

    def test_unverified_has_no_deadline_or_qualifications(self):
        unverified = [
            n for n in self.notices if n["verification_status"] == "UNVERIFIED"
        ]
        self.assertEqual(len(unverified), 3)
        for notice in unverified:
            self.assertIsNone(notice["deadline"])
            self.assertFalse(notice["qualification_verified"])
            self.assertTrue(notice["status_reason"])

    def test_review_state_mix(self):
        states = [p1.compute_review_state(n, self.as_of) for n in self.notices]
        self.assertEqual(states.count("ACTIVE"), 6)
        self.assertEqual(states.count("EXPIRED"), 4)
        self.assertEqual(states.count("NOT_ELIGIBLE"), 1)
        self.assertEqual(states.count("UNVERIFIED"), 3)

    def test_known_deadlines(self):
        self.assertEqual(
            self.by_no["R26BK01762398"]["deadline"], "2026-11-06T10:00:00+09:00"
        )
        self.assertEqual(
            self.by_no["R26BK01787191"]["deadline"], "2026-11-02T18:00:00+09:00"
        )
        self.assertEqual(
            self.by_no["G012605603"]["deadline"], "2026-11-18T14:00:00+09:00"
        )
        self.assertEqual(
            self.by_no["R26BK01762379"]["deadline"], "2026-10-26T16:00:00+09:00"
        )

    def test_qualification_and_document_counts(self):
        self.assertEqual(
            sum(1 for n in self.notices if n.get("qualification_verified")), 8
        )
        self.assertEqual(
            sum(1 for n in self.notices if n.get("documents_verified")), 7
        )


class TestCatalogMatches(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.rows = p1.load_catalog(CATALOG)
        cls.by_cat = p1.build_catalog_index(cls.rows)

    def test_catalog_total(self):
        self.assertEqual(len(self.rows), 4575)

    def test_profile_counts(self):
        self.assertEqual(
            len([r for r in self.rows if r["category_name"] in p1.PROFILES["broadcast_install"]["categories"]]),
            597,
        )
        self.assertEqual(
            len([r for r in self.rows if r["category_name"] in p1.PROFILES["speaker_system"]["categories"]]),
            408,
        )
        self.assertEqual(
            len([r for r in self.rows if r["category_name"] == "디지털 믹서"]), 32
        )

    def test_projector_has_no_catalog_match(self):
        notice = sample_notice(cgi_fit={"level": "LOW", "category_profile": "projector"})
        match = p1.match_notice(notice, self.by_cat)
        self.assertEqual(match["matched_product_count"], 0)
        self.assertEqual(match["sample_products"], "")

    def test_mixer_match_has_examples(self):
        notice = sample_notice(cgi_fit={"level": "HIGH", "category_profile": "mixer"})
        match = p1.match_notice(notice, self.by_cat)
        self.assertEqual(match["matched_product_count"], 32)
        self.assertTrue(match["sample_products"])
        self.assertIn("디지털 믹서:32", match["matched_categories"])


class TestGenerate(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.out_dir = cls.tmp.name
        cls.summary = p1.generate(DATASET, CATALOG, cls.out_dir, AS_OF)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def path(self, name):
        return os.path.join(self.out_dir, name)

    def test_output_files_exist(self):
        for name in [
            "p1_original_verification.csv",
            "p1_original_verification.xlsx",
            "p1_quote_review_items.csv",
            "p1_cgi_product_matches.csv",
            "p1_evidence_manifest.json",
            "p1_summary.json",
            "p1_original_verification.md",
        ]:
            self.assertTrue(os.path.exists(self.path(name)), name)

    def test_summary_counts(self):
        self.assertEqual(self.summary["total_notices"], 14)
        self.assertEqual(self.summary["by_review_state"]["ACTIVE"], 6)
        self.assertEqual(self.summary["by_review_state"]["EXPIRED"], 4)
        self.assertEqual(self.summary["by_review_state"]["NOT_ELIGIBLE"], 1)
        self.assertEqual(self.summary["by_review_state"]["UNVERIFIED"], 3)
        self.assertEqual(self.summary["verified_detail"], 11)
        self.assertEqual(self.summary["verified_deadline"], 11)
        self.assertEqual(self.summary["qualification_verified"], 8)
        self.assertEqual(self.summary["documents_verified"], 7)

    def test_main_csv_has_review_fields(self):
        with open(self.path("p1_original_verification.csv"), encoding="utf-8-sig") as f:
            rows = list(csv.DictReader(f))
        self.assertEqual(len(rows), 14)
        for row in rows:
            self.assertTrue(row["source_url"].startswith("https://"))
            self.assertTrue(row["checked_at"].endswith("+09:00"))
            self.assertIn(row["review_state"], p1.SALES_REVIEW)
            self.assertTrue(row["next_action"])
            if row["verification_status"] == "VERIFIED":
                self.assertTrue(row["cancellation_status"])
        active = [row for row in rows if row["review_state"] == "ACTIVE"]
        self.assertEqual(len(active), 6)
        priorities = [int(row["p1_priority"]) for row in active]
        self.assertEqual(priorities, sorted(priorities))

    def test_unverified_rows_have_no_deadline(self):
        with open(self.path("p1_original_verification.csv"), encoding="utf-8-sig") as f:
            rows = list(csv.DictReader(f))
        for row in rows:
            if row["verification_status"] == "UNVERIFIED":
                self.assertEqual(row["deadline"], "")
                self.assertEqual(row["qualification_verified"], "False")

    def test_quote_review_rows(self):
        with open(self.path("p1_quote_review_items.csv"), encoding="utf-8-sig") as f:
            rows = list(csv.DictReader(f))
        self.assertEqual(len(rows), self.summary["quote_review_items_total"])
        self.assertEqual(len(rows), 45)
        active_ids = {
            entry["bid_no"] for entry in self.summary["active"]
        }
        self.assertTrue(active_ids.issubset({row["bid_no"] for row in rows}))

    def test_manifest_has_per_notice_sources(self):
        with open(self.path("p1_evidence_manifest.json"), encoding="utf-8") as f:
            manifest = json.load(f)
        self.assertEqual(len(manifest["notices"]), 14)
        by_no = {entry["bid_no"]: entry for entry in manifest["notices"]}
        self.assertEqual(
            by_no["R26BK01762398"]["source_url"],
            "https://www.g2b.go.kr/link/PNPE027_01/single/?bidPbancNo=R26BK01762398&bidPbancOrd=000",
        )
        self.assertEqual(
            by_no["G012605603"]["source_url"], "https://srm.kepco.net/index.do"
        )
        for entry in manifest["notices"]:
            self.assertTrue(entry["source_url"])
            if entry["verification_status"] == "UNVERIFIED":
                self.assertIsNone(entry["deadline"])

    def test_xlsx_sheets(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not installed")
        wb = openpyxl.load_workbook(self.path("p1_original_verification.xlsx"))
        self.assertEqual(
            wb.sheetnames,
            [
                "P1_Verification",
                "Quote_Review",
                "Qualifications",
                "Documents",
                "Evidence_Quotes",
                "CGI_Matches",
                "Summary",
            ],
        )
        self.assertEqual(wb["P1_Verification"].max_row, 15)

    def test_markdown_states_policy(self):
        with open(self.path("p1_original_verification.md"), encoding="utf-8") as f:
            text = f.read()
        self.assertIn("실제 원문", text)
        self.assertIn("UNVERIFIED", text)
        self.assertIn("R26BK01762398", text)

    def test_mismatched_state_fails_loudly(self):
        dataset = p1.load_dataset(DATASET)
        dataset["notices"][0]["review_state"] = "EXPIRED"
        as_of = p1.as_kst(p1.parse_dt(AS_OF))
        with self.assertRaises(p1.DatasetError):
            p1.validate_dataset(dataset, as_of)

    def test_unverified_deadline_is_rejected(self):
        dataset = p1.load_dataset(DATASET)
        unverified = next(
            n for n in dataset["notices"] if n["verification_status"] == "UNVERIFIED"
        )
        unverified["deadline"] = "2026-11-01T10:00:00+09:00"
        as_of = p1.as_kst(p1.parse_dt(AS_OF))
        with self.assertRaises(p1.DatasetError):
            p1.validate_dataset(dataset, as_of)

    def test_naive_timestamp_defaults_to_kst(self):
        parsed = p1.parse_dt("2026-10-23T12:00:00")
        self.assertEqual(parsed.utcoffset().total_seconds(), 9 * 3600)


if __name__ == "__main__":
    unittest.main()
