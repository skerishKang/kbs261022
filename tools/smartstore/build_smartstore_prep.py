#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build Naver Smart Store registration preparation files from product_db.xlsx.

This script does NOT perform live registration. It converts the existing CGI
product DB into a normalized preparation file (CSV/XLSX) aligned with the
Naver Smart Store bulk-registration workflow, validates required fields,
reports missing values, and checks duplicate product IDs.

Inputs (relative to repo root):
  data/full_catalog/product_db.xlsx
    - sheet '통합상품'        (canonical product master)
    - sheet '네이버스토어준비' (Naver prep view)
    - sheet '옵션SKU'         (option SKU detail)

Outputs:
  data/smartstore_prep/smartstore_registration_prep.csv   (UTF-8 BOM)
  data/smartstore_prep/smartstore_registration_prep.xlsx
  reports/data/smartstore/missing_fields_report.json
  reports/data/smartstore/duplicate_ids_report.json
  reports/data/smartstore/image_url_mapping.csv
  reports/data/smartstore/summary.json
"""
from __future__ import annotations

import csv
import json
import os
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required: pip install openpyxl") from exc

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DB_PATH = os.path.join(REPO_ROOT, "data", "full_catalog", "product_db.xlsx")
OUT_DIR = os.path.join(REPO_ROOT, "data", "smartstore_prep")
REP_DIR = os.path.join(REPO_ROOT, "reports", "data", "smartstore")

# Naver Smart Store bulk-registration required fields (from official template
# research). These are the fields the seller must provide; values that cannot be
# derived from the CGI DB are left blank and reported as missing.
NAVER_REQUIRED = [
    "카테고리번호",       # 스마트스토어 카테고리 번호 (고객 확정)
    "상품명",
    "판매가",
    "재고수량",
    "대표이미지",
    "상세설명",
    "원산지코드",
    "ASTemplateCode",
    "배송비",
]
NAVER_RECOMMENDED = [
    "브랜드", "제조사", "모델명", "옵션", "검색태그", "판매자코드",
    "전시상태", "KC인증/고시정보", "반품/교환정책",
]

PREP_COLUMNS = [
    "공급처원본ID",
    "네이버상품ID",
    "상품명",
    "브랜드",
    "제조사",
    "모델명",
    "카테고리경로_CGI",
    "스마트스토어카테고리번호",
    "판매가",
    "가격확인필요",
    "재고수량",
    "대표이미지URL",
    "상세이미지수",
    "상세설명",
    "옵션수",
    "옵션SKU요약",
    "원산지코드",
    "수입사",
    "ASTemplateCode",
    "배송비",
    "반품교환정책",
    "KC고시정보",
    "판매자코드",
    "출처URL",
    "등록상태",
    "등록가능여부_DB",
    "사유_DB",
    "누락필수값",
    "준비도",
]


def load_sheets():
    if not os.path.exists(DB_PATH):
        raise FileNotFoundError(DB_PATH)
    wb = openpyxl.load_workbook(DB_PATH, read_only=True, data_only=True)
    required = ["통합상품", "네이버스토어준비", "옵션SKU"]
    for name in required:
        if name not in wb.sheetnames:
            raise ValueError(f"missing sheet: {name}")
    master = list(wb["통합상품"].iter_rows(values_only=True))
    naver = list(wb["네이버스토어준비"].iter_rows(values_only=True))
    options = list(wb["옵션SKU"].iter_rows(values_only=True))
    return master, naver, options


def index_header(rows):
    return {name: i for i, name in enumerate(rows[0])}


def build_option_summary(options_rows):
    """Map 공급처원본ID -> compact option summary string."""
    hdr = index_header(options_rows)
    # option sheet columns: check
    id_col = hdr.get("공급처원본ID")
    if id_col is None:
        # fallback: first column
        id_col = 0
    groups = defaultdict(list)
    for r in options_rows[1:]:
        if not r or r[id_col] is None:
            continue
        pid = str(r[id_col])
        # build a short label from remaining non-null cells
        parts = [str(x) for x in r[1:] if x is not None]
        if parts:
            groups[pid].append(" / ".join(parts[:4]))
    out = {}
    for pid, items in groups.items():
        out[pid] = " | ".join(items[:5])
        if len(items) > 5:
            out[pid] += f" (+{len(items)-5} more)"
    return out


def main_image_ok(url):
    return bool(url) and isinstance(url, str) and url.startswith("http")


def classify_ready(missing_required):
    if not missing_required:
        return "ready"
    # price-only or policy-only gaps are 'partial'; structural gaps 'blocked'
    structural = {"상품명", "대표이미지"}
    if structural & set(missing_required):
        return "blocked"
    return "partial"


def build_prep(master_rows, naver_rows, option_summary):
    m_hdr = index_header(master_rows)
    n_hdr = index_header(naver_rows)

    # map naver prep by source id
    naver_by_id = {}
    for r in naver_rows[1:]:
        if not r:
            continue
        pid = r[n_hdr["공급처원본ID"]]
        if pid is None:
            continue
        naver_by_id[str(pid)] = r

    records = []
    for m in master_rows[1:]:
        if not m:
            continue
        key = m[m_hdr["통합키"]]
        if key is None:
            continue
        pid = str(key)
        n = naver_by_id.get(pid)

        # pull fields (prefer naver prep view, fallback to master)
        name = None
        brand = None
        manufacturer = None
        model = None
        cat_path = None
        price = None
        price_need = None
        main_img = None
        detail_img_count = None
        desc = None
        option_count = None
        naver_product_id = None
        ss_cat_no = None
        reg_status = None
        reg_possible = None
        reason = None

        if n is not None:
            name = n[n_hdr["상품명"]]
            brand = n[n_hdr["브랜드"]]
            model = n[n_hdr["모델명"]]
            cat_path = n[n_hdr["카테고리경로"]]
            price = n[n_hdr["판매가"]]
            price_need = n[n_hdr["가격확인필요"]]
            main_img = n[n_hdr["대표이미지"]]
            detail_img_count = n[n_hdr["상세이미지수"]]
            desc = n[n_hdr["설명"]]
            option_count = n[n_hdr["옵션수"]]
            naver_product_id = n[n_hdr["네이버상품ID"]]
            ss_cat_no = n[n_hdr["스마트스토어카테고리번호"]]
            reg_status = n[n_hdr["등록상태"]]
            reg_possible = n[n_hdr["등록가능여부"]]
            reason = n[n_hdr["사유"]]

        # fallback to master
        if name is None:
            name = m[m_hdr["상품명"]]
        if brand is None:
            brand = m[m_hdr["브랜드"]]
        if model is None:
            model = m[m_hdr["모델명"]]
        if cat_path is None:
            cat_path = m[m_hdr["카테고리경로"]]
        if price is None:
            price = m[m_hdr["공개판매가"]]
        if main_img is None:
            main_img = m[m_hdr["대표이미지"]]
        if detail_img_count is None:
            detail_img_count = m[m_hdr["이미지수"]]
        if option_count is None:
            option_count = m[m_hdr["옵션수"]]
        manufacturer = m[m_hdr["제조사"]]
        source_url = m[m_hdr["출처URL"]]
        seller_code = pid  # candidate seller code = source product id

        # option summary
        opt_summary = option_summary.get(pid, "")
        if option_count and option_count > 0 and not opt_summary:
            opt_summary = f"옵션 {option_count}개 (상세 SKU 미수집)"

        # missing required evaluation (Naver template perspective)
        missing = []
        if not ss_cat_no:
            missing.append("카테고리번호")
        if not name:
            missing.append("상품명")
        if price is None:
            missing.append("판매가")
        if True:  # stock not collected anywhere
            missing.append("재고수량")
        if not main_image_ok(main_img):
            missing.append("대표이미지")
        if not desc:
            missing.append("상세설명")
        # origin code: not in DB
        missing.append("원산지코드")
        # AS template: not in DB
        missing.append("ASTemplateCode")
        # shipping policy: customer
        missing.append("배송비")

        # recommended missing
        rec_missing = []
        if not brand:
            rec_missing.append("브랜드")
        if manufacturer in (None, "", "None"):
            rec_missing.append("제조사")
        if not model:
            rec_missing.append("모델명")
        if not (option_count and option_count > 0):
            pass
        else:
            if not opt_summary:
                rec_missing.append("옵션")
        if not seller_code:
            rec_missing.append("판매자코드")

        ready = classify_ready(missing)

        records.append({
            "공급처원본ID": pid,
            "네이버상품ID": naver_product_id,
            "상품명": name,
            "브랜드": brand,
            "제조사": manufacturer,
            "모델명": model,
            "카테고리경로_CGI": cat_path,
            "스마트스토어카테고리번호": ss_cat_no,
            "판매가": price,
            "가격확인필요": price_need,
            "재고수량": None,
            "대표이미지URL": main_img,
            "상세이미지수": detail_img_count,
            "상세설명": desc,
            "옵션수": option_count,
            "옵션SKU요약": opt_summary,
            "원산지코드": None,
            "수입사": None,
            "ASTemplateCode": None,
            "배송비": None,
            "반품교환정책": None,
            "KC고시정보": None,
            "판매자코드": seller_code,
            "출처URL": source_url,
            "등록상태": reg_status,
            "등록가능여부_DB": reg_possible,
            "사유_DB": reason,
            "누락필수값": ";".join(missing),
            "준비도": ready,
        })

    return records


def write_csv(records, path):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=PREP_COLUMNS)
        w.writeheader()
        for r in records:
            w.writerow(r)


def write_xlsx(records, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "스마트스토어등록준비"
    ws.append(PREP_COLUMNS)
    for r in records:
        ws.append([r[c] for c in PREP_COLUMNS])
    # header styling
    from openpyxl.styles import Font, PatternFill
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    ws.freeze_panes = "C2"
    wb.save(path)


def write_image_mapping(records, path):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["공급처원본ID", "대표이미지URL", "상세이미지수", "이미지상태"])
        for r in records:
            url = r["대표이미지URL"]
            status = "OK" if (url and isinstance(url, str) and url.startswith("http")) else "MISSING"
            w.writerow([r["공급처원본ID"], url, r["상세이미지수"], status])


def duplicate_report(master_rows):
    hdr = index_header(master_rows)
    key_col = hdr["통합키"]
    seen = defaultdict(list)
    for i, r in enumerate(master_rows[1:], start=2):
        if not r:
            continue
        k = r[key_col]
        if k is None:
            continue
        seen[str(k)].append(i)
    dups = {k: v for k, v in seen.items() if len(v) > 1}
    return {
        "total_keys": len(seen),
        "duplicate_key_count": len(dups),
        "duplicate_rows_extra": sum(len(v) - 1 for v in dups.values()),
        "samples": {k: v for k, v in list(dups.items())[:20]},
    }


def missing_fields_report(records):
    # count missing per required field across all records
    req_counter = Counter()
    ready_counts = Counter()
    for r in records:
        for field in r["누락필수값"].split(";"):
            if field:
                req_counter[field] += 1
        ready_counts[r["준비도"]] += 1
    # DB-level price availability
    price_missing = sum(1 for r in records if r["판매가"] is None)
    return {
        "total_products": len(records),
        "required_field_missing_counts": dict(sorted(req_counter.items())),
        "ready_classification_counts": dict(ready_counts),
        "price_available_count": len(records) - price_missing,
        "price_missing_count": price_missing,
        "image_ok_count": sum(1 for r in records if r["대표이미지URL"] and str(r["대표이미지URL"]).startswith("http")),
        "description_present_count": sum(1 for r in records if r["상세설명"]),
        "option_products_count": sum(1 for r in records if (r["옵션수"] or 0) > 0),
    }


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    os.makedirs(REP_DIR, exist_ok=True)
    master_rows, naver_rows, option_rows = load_sheets()
    option_summary = build_option_summary(option_rows)
    records = build_prep(master_rows, naver_rows, option_summary)

    csv_path = os.path.join(OUT_DIR, "smartstore_registration_prep.csv")
    xlsx_path = os.path.join(OUT_DIR, "smartstore_registration_prep.xlsx")
    img_path = os.path.join(REP_DIR, "image_url_mapping.csv")
    missing_path = os.path.join(REP_DIR, "missing_fields_report.json")
    dup_path = os.path.join(REP_DIR, "duplicate_ids_report.json")
    summary_path = os.path.join(REP_DIR, "summary.json")

    write_csv(records, csv_path)
    write_xlsx(records, xlsx_path)
    write_image_mapping(records, img_path)

    mr = missing_fields_report(records)
    with open(missing_path, "w", encoding="utf-8") as f:
        json.dump(mr, f, ensure_ascii=False, indent=2)

    dr = duplicate_report(master_rows)
    with open(dup_path, "w", encoding="utf-8") as f:
        json.dump(dr, f, ensure_ascii=False, indent=2)

    summary = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_db": "data/full_catalog/product_db.xlsx",
        "total_products": len(records),
        "ready_for_price_count": mr["price_available_count"],
        "blocked_price_count": mr["price_missing_count"],
        "prep_ready_count": sum(1 for r in records if r["준비도"] == "ready"),
        "prep_partial_count": sum(1 for r in records if r["준비도"] == "partial"),
        "prep_blocked_count": sum(1 for r in records if r["준비도"] == "blocked"),
        "duplicate_keys": dr["duplicate_key_count"],
        "outputs": {
            "csv": "data/smartstore_prep/smartstore_registration_prep.csv",
            "xlsx": "data/smartstore_prep/smartstore_registration_prep.xlsx",
            "missing_fields_report": "reports/data/smartstore/missing_fields_report.json",
            "duplicate_ids_report": "reports/data/smartstore/duplicate_ids_report.json",
            "image_url_mapping": "reports/data/smartstore/image_url_mapping.csv",
        },
        "live_registration_performed": False,
        "store_approval_verified": False,
    }
    with open(summary_path, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
