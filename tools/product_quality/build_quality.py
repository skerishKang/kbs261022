"""이슈 #43 상품 DB 품질 개선 산출물 빌더 (Stage 1 + Stage 2).

실행 (저장소 루트):

    python tools\\product_quality\\build_quality.py

읽기 전용 입력
    - ``site/data/products.json``           원본 마스터 4,575건 (덮어쓰지 않음)
    - ``data/full_catalog/product_db.xlsx`` 이슈 #8 정본 — 중복 판정 재사용

Stage 1: 원본 스펙행과 confirmed 중복 승계만 자동 보완한다.
Stage 2 (``verify_review``): 검토 후보를 이미 수집된 모델명 어휘로 교차 실측해
독립 원본 ID 2건 이상으로 실증된 것만 auto_confirmed 로 승격한다 (--no-stage2 로 끄기).

쓰기 위치: ``data/product_quality/`` (신규 영역. LOCAL2/LOCAL3 파일과 충돌하지 않음)
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import sys
from typing import Any, Mapping, Sequence

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from tools.product_quality import quality as Q  # noqa: E402
from tools.product_quality import verify_review as verify  # noqa: E402

OUT_DIR = "data/product_quality"
FILES = {
    "patch_candidates": "product_quality_patch_candidates.csv",
    "duplicate_candidates": "product_duplicate_candidates.csv",
    "improved": "product_quality_improved.json",
    "summary": "product_quality_summary.json",
    "review": "product_quality_review.xlsx",
    "stage2": "product_quality_stage2_verification.csv",
}

# 정본 XLSX 중복후보 시트 → (좌측, 우측) 열 이름
SHEET_DUPLICATE = "중복후보"
COLS = {
    "verdict": "판정",
    "auto": "자동병합가능",
    "left_site": "왼쪽공급처",
    "left_id": "왼쪽원본ID",
    "left_name": "왼쪽상품명",
    "left_brand": "왼쪽브랜드",
    "left_model": "왼쪽모델명",
    "right_site": "오른쪽공급처",
    "right_id": "오른쪽원본ID",
    "right_name": "오른쪽상품명",
    "right_brand": "오른쪽브랜드",
    "right_model": "오른쪽모델명",
    "phrase": "공통모델표현",
    "similarity": "이름유사도",
    "conflict": "변형충돌",
    "brand_ok": "브랜드값일치",
    "reason": "사유",
}


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _key(site_label: str, product_id: Any) -> str:
    site = Q.SITE_BY_LABEL.get(_clean(site_label), _clean(site_label))
    return f"{site}:{_clean(product_id)}"


def load_master(path: str) -> list[dict[str, Any]]:
    with io.open(path, "r", encoding="utf-8") as handle:
        records = json.load(handle)
    if not isinstance(records, list):
        raise SystemExit(f"{path}: 원본 마스터는 리스트여야 합니다")
    return records


def load_duplicate_sheet(path: str) -> tuple[list[Q.DuplicateCandidate], list[Mapping[str, Any]]]:
    """정본의 중복 판정을 재사용한다 (재계산·재수집 없음)."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[SHEET_DUPLICATE]
    rows = sheet.iter_rows(values_only=True)
    header = [_clean(h) for h in next(rows)]
    index = {name: header.index(name) for name in COLS.values() if name in header}
    missing = [name for name in COLS.values() if name not in index]
    if missing:
        raise SystemExit(f"{path}:{SHEET_DUPLICATE} 필요한 열 없음: {missing}")

    master_pairs: list[Q.DuplicateCandidate] = []
    confirmed_pairs: list[Mapping[str, Any]] = []
    for raw in rows:
        if raw is None or all(c is None or _clean(c) == "" for c in raw):
            continue
        get = lambda name: _clean(raw[index[name]])  # noqa: E731
        left_key = _key(get(COLS["left_site"]), get(COLS["left_id"]))
        right_key = _key(get(COLS["right_site"]), get(COLS["right_id"]))
        verdict = get(COLS["verdict"])
        auto = get(COLS["auto"]) == "예"
        master_pairs.append(
            Q.DuplicateCandidate(
                verdict=verdict,
                auto_mergeable=auto,
                left_key=left_key,
                left_name=get(COLS["left_name"]),
                left_brand=get(COLS["left_brand"]),
                left_model=get(COLS["left_model"]),
                right_key=right_key,
                right_name=get(COLS["right_name"]),
                right_brand=get(COLS["right_brand"]),
                right_model=get(COLS["right_model"]),
                shared_phrase=get(COLS["phrase"]),
                name_similarity=get(COLS["similarity"]),
                variant_conflict=get(COLS["conflict"]),
                brand_agreement=get(COLS["brand_ok"]),
                reason=get(COLS["reason"]),
                origin="issue8_master",
            )
        )
        if verdict == "confirmed":
            confirmed_pairs.append(
                {"left_key": left_key, "right_key": right_key, "shared_phrase": get(COLS["phrase"])}
            )
    workbook.close()
    return master_pairs, confirmed_pairs


def patch_rows_with_metrics(
    patches: Sequence[Q.Patch],
    verification: Mapping[tuple[str, str], Any],
) -> list[dict[str, Any]]:
    """패치 행에 Stage 2 실측 지표(실측 횟수·독립 ID 수·일치·중복 등급·우선순위)를 병합한다."""
    rows: list[dict[str, Any]] = []
    for patch in patches:
        row = patch.to_row()
        ver = verification.get((patch.product_key, patch.field))
        row.update(ver.to_columns() if ver else {})
        rows.append(row)
    return rows


def write_rows_csv(rows: Sequence[Mapping[str, Any]], path: str) -> int:
    header = list(rows[0].keys())
    with io.open(path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=header)
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


def write_patch_csv(patches: Sequence[Q.Patch], path: str,
                    verification: Mapping[tuple[str, str], Any]) -> int:
    return write_rows_csv(patch_rows_with_metrics(patches, verification), path)


def write_stage2_csv(
    patches: Sequence[Q.Patch],
    verification: Mapping[tuple[str, str], Any],
    path: str,
) -> int:
    """Stage 2가 검증한 행(승격 + 우선순위 부여 검토)만 별도 산출물로 남긴다."""
    rows = []
    for patch in patches:
        ver = verification.get((patch.product_key, patch.field))
        if ver is None or (not ver.promoted and not ver.priority):
            continue
        row = patch.to_row()
        row.update(ver.to_columns())
        row["Stage2승격"] = "예" if ver.promoted else "아니오"
        rows.append(row)
    rows.sort(key=lambda r: (r.get("검토우선순위", "9"), r["필드"], r["통합키"]))
    return write_rows_csv(rows, path)


def write_duplicate_csv(duplicates: Sequence[Q.DuplicateCandidate], path: str) -> int:
    rows = [d.to_row() for d in duplicates]
    header = list(rows[0].keys())
    with io.open(path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=header)
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


def write_json(payload: Any, path: str, compact: bool = False) -> None:
    with io.open(path, "w", encoding="utf-8") as handle:
        if compact:
            json.dump(payload, handle, ensure_ascii=False, separators=(",", ":"))
        else:
            json.dump(payload, handle, ensure_ascii=False, indent=1)
        handle.write("\n")


def write_review_xlsx(result: Q.QualityResult, summary: Mapping[str, Any], path: str,
                      verification: Mapping[tuple[str, str], Any],
                      promoted: Sequence[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    bold = Font(bold=True)

    def sheet(title: str, rows: Sequence[Mapping[str, Any]], note: str = ""):
        ws = workbook.create_sheet(title=title)
        if note:
            ws.append([note])
            ws.cell(row=1, column=1).font = bold
        if not rows:
            ws.append(["(해당 없음)"])
            return ws
        header = list(rows[0].keys())
        ws.append(header)
        for cell in ws[ws.max_row]:
            cell.font = bold
        for row in rows:
            ws.append([_cell(row.get(name)) for name in header])
        ws.freeze_panes = ws["A3" if note else "A2"]
        for i, name in enumerate(header, start=1):
            width = min(46, max(len(str(name)) + 2, 12))
            ws.column_dimensions[get_column_letter(i)].width = width
        return ws

    def _cell(value: Any) -> Any:
        if isinstance(value, list) and value and isinstance(value[0], dict):
            return f"{len(value)}건 (Stage2승격 시트 참조)"
        if isinstance(value, (list, tuple)):
            return ", ".join(str(v) for v in value)
        if isinstance(value, dict):
            return json.dumps(value, ensure_ascii=False)
        if isinstance(value, bool):
            return "예" if value else "아니오"
        return value

    keyed = patch_rows_with_metrics(result.patches, verification)
    auto_rows = [r for r, p in zip(keyed, result.patches)
                 if p.classification == Q.CLASS_AUTO and p.proposed != p.before]
    review_rows = sorted((r for r, p in zip(keyed, result.patches) if p.classification == Q.CLASS_REVIEW),
                         key=lambda r: (str(r.get("검토우선순위") or "9"), r["필드"], r["통합키"]))
    none_rows = [r for r, p in zip(keyed, result.patches) if p.classification == Q.CLASS_NONE]
    promoted_rows = [r for r, p in zip(keyed, result.patches)
                     if (p.product_key, p.field) and p.classification == Q.CLASS_AUTO
                     and p.evidence_type == "multi_source_attestation"]
    field_rows = field_status_rows(result)
    contract_rows = local2_contract_rows(summary)

    del workbook["Sheet"]
    sheet("요약", [{k: _cell(v) for k, v in dict(flatten(summary)).items()}],
          "이슈 #43 상품 DB 품질 개선 Stage 2 — 보완 전후 요약 (원본 4,575건 보존)")
    sheet("Stage2승격", promoted_rows,
          f"Stage 2 자동 승격 {len(promoted_rows)}건: 같은 브랜드 모델명 스펙행 독립 원본 ID 2건 이상 실측")
    sheet("자동보완", auto_rows, "원천 근거가 유일하거나 Stage 2 교차 실측을 통과. LOCAL2는 즉시 사용 가능")
    sheet("검토후보", review_rows,
          "근거가 약하거나 복수. 우선순위 1→4 순 정렬(실측횟수·독립ID수·영문/파생 일치·중복 등급 포함)")
    sheet("근거없음", none_rows, "원천에 표기가 없어 공백 유지. 임의 생성 금지")
    sheet("중복후보", [d.to_row() for d in result.duplicates],
          "이슈 #8 정본 판정 재사용 + 같은 사이트 완전 일치. 원본 보존, 병합 0건")
    sheet("필드현황", field_rows, "필드별 누락 전후 및 공급처 분해")
    sheet("LOCAL2계약", contract_rows, "product_quality_improved.json 열 의미")

    workbook.save(path)


def flatten(summary: Mapping[str, Any], prefix: str = "") -> list[tuple[str, Any]]:
    out: list[tuple[str, Any]] = []
    for key, value in summary.items():
        name = f"{prefix}{key}"
        if isinstance(value, dict):
            out.extend(flatten(value, name + "."))
        else:
            out.append((name, value))
    return out


def field_status_rows(result: Q.QualityResult) -> list[dict[str, Any]]:
    by_field: dict[tuple[str, str], list[Q.Patch]] = {}
    for patch in result.patches:
        by_field.setdefault((patch.field, patch.site), []).append(patch)
    counts: dict[tuple[str, str], dict[str, int]] = {}
    for (field_name, site), patches in by_field.items():
        actionable = [p for p in patches if p.actionable]
        counts[(field_name, site)] = {
            "자동보완": sum(1 for p in actionable if p.classification == Q.CLASS_AUTO),
            "검토후보": sum(1 for p in actionable if p.classification == Q.CLASS_REVIEW),
            "근거없음": sum(1 for p in actionable if p.classification == Q.CLASS_NONE),
            "이미원본값": len(patches) - len(actionable),
            "자리표시자": sum(1 for p in patches
                              if p.evidence_type in ("placeholder_existing_value", "placeholder_replacement")),
        }
    rows: list[dict[str, Any]] = []
    for field_name in ("model", "manufacturer", "brand"):
        for site in ("samasound.co.kr", "samamusic.co.kr"):
            subset = [r for r in result.improved if r["source_site"] == site]
            if field_name == "brand":
                missing_before = sum(1 for r in subset if not _clean(r.get("brand")))
                missing_after = missing_before
            else:
                missing_before = sum(1 for r in subset if not r[f"{field_name}_before"])
                missing_after = sum(1 for r in subset if not r[f"{field_name}_after"])
            stat = counts.get((field_name, site), {})
            rows.append(
                {
                    "필드": field_name,
                    "공급처": site,
                    "상품수": len(subset),
                    "보완전누락": missing_before,
                    "보완후누락": missing_after,
                    **stat,
                }
            )
    return rows


def local2_contract_rows(summary: Mapping[str, Any]) -> list[dict[str, Any]]:
    stage2 = summary.get("stage2") or {}
    contract_version = summary.get("schema_version", Q.SCHEMA_VERSION)
    promoted = stage2.get("auto_promoted_model", 0)
    return [
        {"항목": "파일", "내용": "data/product_quality/product_quality_improved.json"},
        {"항목": "schema_version", "내용": f"{contract_version} (Stage 2 승격 반영본)"},
        {"항목": "성격", "내용": "패치 레이어. 원본 site/data/products.json 을 덮어쓰지 않는다"},
        {"항목": "적용 방법", "내용": "products.json 의 id 와 improved.products[].id 를 조인한 뒤 "
                                        "manufacturer_after / model_after 를 우선 사용"},
        {"항목": "즉시 사용 가능", "내용": "classification=auto_confirmed 항목 (총 {}건 레코드)".format(
            summary["auto_corrected_records"])},
        {"항목": "Stage 2 승격", "내용": "모델명 {}건이 Stage 2에서 auto_confirmed 로 승격되어 위 계약으로 "
                                        "그대로 전달된다. 승격 근거는 product_quality_stage2_verification.csv"
                                        " 및 summary.stage2.promotions".format(promoted)},
        {"항목": "검토 우선행", "내용": "verification.model.priority=1 행은 사람이 먼저 확인할 다음 후보. "
                                          "단 승격(promoted)이 아니면 적용 금지"},
        {"항목": "반영 금지", "내용": "classification=review_required / no_evidence 항목. "
                                       "product_quality_review.xlsx 검토후보 시트 확정 전까지 원본값 유지"},
        {"항목": "가격·재고", "내용": "price_original / price_state_original 은 원본 전달값뿐이며 "
                                        "새로 만든 값이 없다"},
        {"항목": "중복", "내용": "중복 후보는 표시만 했다. 병합 0건, 원본 4,575건 그대로"},
        {"항목": "Yamaha", "내용": "673건 재수집 없음. site/data/yamaha_catalog.json 은 이 작업에서 읽지 않았다"},
    ]


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="이슈 #43 상품 DB 품질 개선 산출물 생성")
    parser.add_argument("--master", default=Q.MASTER_PATH)
    parser.add_argument("--authoritative", default=Q.AUTHORITATIVE_XLSX)
    parser.add_argument("--out-dir", default=OUT_DIR)
    parser.add_argument("--no-stage2", action="store_true",
                        help="Stage 1 산출물만 생성 (회귀 비교용)")
    parser.add_argument("--quiet", action="store_true", help="요약 JSON 출력을 생략한다")
    args = parser.parse_args(argv)

    os.makedirs(args.out_dir, exist_ok=True)
    records = load_master(args.master)
    master_pairs, confirmed_pairs = load_duplicate_sheet(args.authoritative)
    result = Q.build_result(records, confirmed_pairs, master_pairs)

    stage2 = verify.Stage2Result()
    stage1_summary = json.loads(json.dumps(result.summary))
    if not args.no_stage2:
        # Stage 1 수치를 먼저 확보한 뒤 승격을 반영해 재집계한다 (전후 수치 비교용).
        stage2 = verify.apply(result, records)
        result.summary = Q.summarize(result)
        result.summary["schema_version"] = verify.STAGE2_VERSION
        result.summary["stage1_before_stage2"] = stage1_summary
        result.summary["stage2"] = {
            **stage2.stats,
            "priority_labels": dict(verify.PRIORITY_LABELS),
            "promotions": stage2.promotions,
        }

    actionable = [p for p in result.patches if p.actionable]
    patch_rows = write_patch_csv(actionable, os.path.join(args.out_dir, FILES["patch_candidates"]),
                                 stage2.verification)
    dup_rows = write_duplicate_csv(result.duplicates, os.path.join(args.out_dir, FILES["duplicate_candidates"]))
    stage2_rows = write_stage2_csv(result.patches, stage2.verification,
                                   os.path.join(args.out_dir, FILES["stage2"])) if stage2.stats else 0
    improved_payload = {
        "schema_version": result.summary["schema_version"],
        "generated_at": result.generated_at,
        "issue": 43,
        "layer": "patch",
        "base_master": "site/data/products.json",
        "original_products": result.master_count,
        "stage2": {
            "rule": stage2.stats.get("rule", ""),
            "auto_promoted_model": stage2.stats.get("auto_promoted_model", 0),
            "model_missing_after_stage1": stage2.stats.get("model_missing_after_stage1"),
            "model_missing_after_stage2": stage2.stats.get("model_missing_after_stage2"),
        },
        "products": result.improved,
    }
    write_json(improved_payload, os.path.join(args.out_dir, FILES["improved"]), compact=True)
    write_json(result.summary, os.path.join(args.out_dir, FILES["summary"]))
    write_review_xlsx(result, result.summary, os.path.join(args.out_dir, FILES["review"]),
                      stage2.verification, stage2.promoted)

    if args.quiet:
        return 0
    console_summary = json.loads(json.dumps(result.summary, default=str))
    console_summary.get("stage2", {}).pop("promotions", None)
    print(json.dumps(console_summary, ensure_ascii=False, indent=1, default=str))
    print(f"patch_rows={patch_rows} duplicate_rows={dup_rows} stage2_rows={stage2_rows} "
          f"promoted={len(stage2.promoted)} out_dir={args.out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
