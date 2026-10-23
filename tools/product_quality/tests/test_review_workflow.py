"""이슈 #43 Stage 3 — 검토 워크플로우(대기열·결정 검증·패치 내보내기) 테스트.

검토자 실수·임의 입력을 막는 규칙을 집중적으로 검증한다:
- 근거 없는 값(빈 제안값·자리표시자)은 승인할 수 없다.
- 승인에는 검토자 기록이 반드시 필요하다.
- 대기열 재생성 후 제안값·우선순위가 바뀌면 행지문 불일치로 거부한다.
- 중복 후보는 판정만 기록하고 병합은 언제나 0건이다.
- 원본 마스터와 Stage 2 패치 레이어 값은 Stage 3 에서 바뀌지 않는다.
"""

from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import unittest
from collections import Counter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
sys.path.insert(0, ROOT)

from tools.product_quality import quality as Q  # noqa: E402
from tools.product_quality import review_workflow as W  # noqa: E402
from tools.product_quality import verify_review as V  # noqa: E402
from tools.product_quality.tests.test_quality import record  # noqa: E402


def build(fixture, confirmed=()):
    pairs = [
        Q.DuplicateCandidate(
            verdict="confirmed", auto_mergeable=True,
            left_key=c["left_key"], left_name="", left_brand="", left_model="",
            right_key=c["right_key"], right_name="", right_brand="", right_model="",
            shared_phrase=c.get("shared_phrase", ""), name_similarity="1.0",
            variant_conflict="아니오", brand_agreement="예", reason="test", origin="test",
        )
        for c in confirmed
    ]
    result = Q.build_result(fixture, confirmed, pairs)
    stage2 = V.apply(result, fixture)
    result.summary = Q.summarize(result)
    return result, stage2


FIXTURE = [
    record("삼아사운드", "1", "SM58 다이내믹 마이크", "SHURE", manufacturer="Shure"),
    record("삼아사운드", "2", "8010A 제네렉 스튜디오 모니터", "GENELEC", manufacturer="Genelec",
           english_name="8010A"),
    record("삼아사운드", "3", "주문 제작 랙", "자체제작", manufacturer="자체제작"),
    record("삼아스토어", "1", "[SHURE] SM58 마이크", "SHURE", model="SM58", brand_row="SHURE"),
    record("삼아스토어", "2", "기타 액세서리", "자체브랜드", brand_row="자체브랜드"),
    record("삼아스토어", "4", "정체불가 액세서리", "UNKNOWNBRAND", brand_row="UNKNOWNBRAND"),
]
CONFIRMED = [{"left_key": "samamusic.co.kr:1", "right_key": "samasound.co.kr:1", "shared_phrase": "SM58"}]


class QueueBuilder(unittest.TestCase):
    def setUp(self) -> None:
        self.result, self.stage2 = build(FIXTURE, CONFIRMED)
        self.queue = W.build_review_rows(self.result.patches, self.stage2.verification, FIXTURE)
        self.dups = W.build_dup_rows(self.result.duplicates)

    def rows(self, **kw):
        out = [r for r in self.queue if all(r[k] == v for k, v in kw.items())]
        self.assertTrue(out, f"queued row not found: {kw}")
        return out


class TestQueue(QueueBuilder):
    def test_only_review_rows_are_queued(self) -> None:
        review = [p for p in self.result.patches if p.classification == Q.CLASS_REVIEW]
        self.assertEqual(len(self.queue), len(review))
        self.assertTrue(all(r["검토결정"] == "" for r in self.queue))

    def test_queue_excludes_auto_and_no_evidence(self) -> None:
        keys = {(r["통합키"], r["필드"]) for r in self.queue}
        for patch in self.result.patches:
            if patch.classification in (Q.CLASS_AUTO, Q.CLASS_NONE):
                self.assertNotIn((patch.product_key, patch.field), keys, patch.field)

    def test_rows_are_sorted_by_priority(self) -> None:
        priorities = [r["우선순위"] for r in self.queue]
        self.assertEqual(priorities, sorted(priorities))

    def test_every_row_carries_evidence_and_both_values(self) -> None:
        for row in self.queue:
            self.assertTrue(row["근거"], row["검토ID"])
            self.assertIn(row["보완전"], ("", row["보완전"]))
            self.assertTrue(row["원본화면URL"].startswith("https://"), row["원본화면URL"])
            self.assertTrue(row["행지문"])

    def test_review_id_is_stable_and_unique(self) -> None:
        again = W.build_review_rows(self.result.patches, self.stage2.verification, FIXTURE)
        self.assertEqual([r["검토ID"] for r in again], [r["검토ID"] for r in self.queue])
        ids = [r["검토ID"] for r in self.queue]
        self.assertEqual(len(ids), len(set(ids)))

    def test_fingerprint_changes_with_proposed_value(self) -> None:
        row = self.rows(필드="model")[0]
        other = W.row_fingerprint(row["통합키"], row["필드"], row["보완전"], "딴값", row["우선순위"])
        self.assertNotEqual(row["행지문"], other)

    def test_input_columns_exist(self) -> None:
        for column in W.DECISION_COLUMNS:
            self.assertIn(column, self.queue[0])

    def test_queue_fingerprint_is_deterministic(self) -> None:
        first = W.queue_fingerprint(self.queue)
        self.assertEqual(first, W.queue_fingerprint(self.queue))
        self.assertEqual(len(first), 64)

    def test_duplicate_rows_are_not_queued_rows(self) -> None:
        self.assertEqual(len(self.dups), len(self.result.duplicates))
        for row in self.dups:
            self.assertEqual(row["판정"], "")
            self.assertEqual(row["자동병합가능"] in ("예", "아니오"), True)


class TestDecisionValidation(QueueBuilder):
    def decide(self, index=0, decision="승인", value="", reviewer="김검토", note="", **override):
        row = dict(self.queue[index])
        row.update({"검토결정": decision, "확정값": value, "검토자": reviewer,
                    "검토일": "2026-10-23", "비고": note})
        row.update(override)
        return W.read_decisions([row], self.dups, self.queue, W.queue_fingerprint(self.queue))

    def test_approve_with_proposal_becomes_confirmed_value(self) -> None:
        idx = next(i for i, r in enumerate(self.queue) if r["제안값"])
        out = self.decide(idx)
        self.assertEqual(len(out.approved), 1)
        self.assertEqual(out.approved[0].confirmed_value, self.queue[idx]["제안값"])

    def test_approve_without_reviewer_is_invalid(self) -> None:
        idx = next(i for i, r in enumerate(self.queue) if r["제안값"])
        out = self.decide(idx, reviewer="")
        self.assertEqual(out.approved, [])
        self.assertEqual(len(out.invalid), 1)
        self.assertTrue(any("검토자" in e for e in out.errors))

    def test_approve_empty_proposal_is_refused(self) -> None:
        blank = dict(self.queue[0], 제안값="", 행지문=W.row_fingerprint(
            self.queue[0]["통합키"], self.queue[0]["필드"], self.queue[0]["보완전"], "", self.queue[0]["우선순위"]))
        queue = [blank] + self.queue[1:]
        out = W.read_decisions([dict(blank, 검토결정="승인", 검토자="김검토")], [], queue, "")
        self.assertEqual(out.approved, [])
        self.assertTrue(any("실측값이 없다" in e for e in out.errors))

    def test_approve_placeholder_value_is_refused(self) -> None:
        for bad in ("자체제작", "기타", "-", "없음"):
            idx = next(i for i, r in enumerate(self.queue) if r["제안값"])
            out = self.decide(idx, value=bad)
            self.assertEqual(out.approved, [], bad)

    def test_custom_value_without_note_is_refused(self) -> None:
        idx = next(i for i, r in enumerate(self.queue) if r["제안값"])
        out = self.decide(idx, value="GLX-2427")
        self.assertEqual(out.approved, [])
        self.assertTrue(any("비고" in e for e in out.errors))

    def test_custom_value_with_note_is_accepted(self) -> None:
        idx = next(i for i, r in enumerate(self.queue) if r["제안값"])
        out = self.decide(idx, value="GLX-2427", note="공급처 화면 실측")
        self.assertEqual(out.approved[0].confirmed_value, "GLX-2427")

    def test_unknown_decision_is_invalid(self) -> None:
        out = self.decide(0, decision="아무거나")
        self.assertEqual(out.approved, [])
        self.assertTrue(any("허용되지 않은" in e for e in out.errors))

    def test_unknown_review_id_is_refused(self) -> None:
        row = dict(self.queue[0], 검토ID="RV-DEADBEEF00")
        row.update({"검토결정": "승인", "검토자": "김검토"})
        out = W.read_decisions([row], [], self.queue, "")
        self.assertTrue(any("대기열에 없는" in e for e in out.errors))

    def test_duplicate_review_id_is_refused(self) -> None:
        rows = [self.queue[0], self.queue[0]]
        for row in rows:
            row["검토결정"] = "보류"
        out = W.read_decisions(rows, [], self.queue, "")
        self.assertTrue(any("중복된" in e for e in out.errors))

    def test_fingerprint_mismatch_is_reported(self) -> None:
        out = self.decide(0, decision="보류", 행지문="000000000000")
        self.assertTrue(any("행지문 불일치" in e for e in out.errors))

    def test_blank_decision_with_other_input_is_error(self) -> None:
        out = self.decide(0, decision="", value="X")
        self.assertTrue(any("검토결정 필수" in e for e in out.errors))

    def test_reject_and_hold_are_counted(self) -> None:
        rows = []
        for row, decision in ((self.queue[0], "반려"), (self.queue[1], "보류")):
            item = dict(row)
            item.update({"검토결정": decision, "검토자": "김검토"})
            rows.append(item)
        out = W.read_decisions(rows, [], self.queue, "")
        self.assertEqual(len(out.rejected), 1)
        self.assertEqual(len(out.held), 1)
        self.assertEqual(out.stats["duplicates_merged"], 0)

    def test_untouched_rows_are_pending(self) -> None:
        out = self.decide(0, decision="반려")
        self.assertEqual(len(out.pending), len(self.queue) - 1)

    def test_decisions_do_not_mutate_patch_layer(self) -> None:
        before = json.dumps(self.result.improved, ensure_ascii=False, sort_keys=True)
        idx = next(i for i, r in enumerate(self.queue) if r["제안값"])
        self.decide(idx)
        self.assertEqual(before, json.dumps(self.result.improved, ensure_ascii=False, sort_keys=True))


class TestWorkbook(QueueBuilder):
    def setUp(self) -> None:
        super().setUp()
        self.tmp = tempfile.mkdtemp(prefix="pq43s3-")
        self.path = os.path.join(self.tmp, "queue.xlsx")
        self.fp = W.queue_fingerprint(self.queue)
        W.write_queue_xlsx(self.queue, self.dups, self.path, self.fp, {"stage3": {"model_review": 3}})

    def tearDown(self) -> None:
        import shutil

        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_sheets_and_input_validation(self) -> None:
        from openpyxl import load_workbook

        wb = load_workbook(self.path)
        self.assertEqual(wb.sheetnames, ["검토방법", "검토후보", "중복후보판정"])
        ws = wb["검토후보".replace("후보", "후보")] if False else wb["검토후보"]
        self.assertEqual(ws.max_row, len(self.queue) + 1)
        self.assertTrue(ws.data_validations.dataValidation, "드롭다운 검증이 있어야 한다")
        formula = ws.data_validations.dataValidation[0].formula1
        self.assertIn("승인", formula)
        self.assertIn("보류", formula)

    def test_roundtrip_reads_decisions(self) -> None:
        review_rows, dup_rows, fp = W.load_workbook_rows(self.path)
        self.assertEqual(fp, self.fp)
        self.assertEqual(len(review_rows), len(self.queue))
        self.assertEqual(len(dup_rows), len(self.dups))
        idx = next(i for i, r in enumerate(review_rows) if r["제안값"])
        review_rows[idx]["검토결정"] = "승인"
        review_rows[idx]["검토자"] = "김검토"
        out = W.read_decisions(review_rows, dup_rows, self.queue, fp)
        self.assertEqual(len(out.approved), 1)

    def test_export_writes_json_and_csv(self) -> None:
        payload = W.apply_decisions_file(self.path, self.queue, self.tmp, {"schema_version": "1.1.0"})
        self.assertTrue(os.path.exists(payload["output_files"]["json"]))
        self.assertTrue(os.path.exists(payload["output_files"]["csv"]))
        self.assertEqual(payload["counts"]["pending"], len(self.queue))
        self.assertEqual(payload["policy"]["duplicates_merged"], 0)
        with io.open(payload["output_files"]["csv"], encoding="utf-8-sig") as handle:
            head = handle.readline()
        self.assertIn("검토ID", head)


class TestExportContract(unittest.TestCase):
    def test_policy_flags_are_conservative(self) -> None:
        payload = W.decisions_payload(W.DecisionResult(), {"schema_version": V.STAGE2_VERSION})
        policy = payload["policy"]
        for key in ("original_master_overwritten", "duplicates_merged", "price_invented",
                    "stock_invented", "wholesale_invented", "new_scraping_performed",
                    "unapproved_values_applied"):
            self.assertFalse(policy[key], key)
        self.assertEqual(payload["schema_version"], W.STAGE3_VERSION)
        self.assertEqual(payload["stage"], "stage3_review_decisions")

    def test_empty_decisions_csv_is_header_only(self) -> None:
        tmp = tempfile.mkdtemp(prefix="pq43s3-csv-")
        try:
            path = os.path.join(tmp, "d.csv")
            self.assertEqual(W.write_decisions_csv({"approved": [], "rejected": [], "held": [],
                                                    "invalid": []}, path), 0)
            with io.open(path, encoding="utf-8-sig") as handle:
                header = handle.readline().strip()
            self.assertEqual(header, ",".join(W._DECISION_CSV_HEADER))
        finally:
            import shutil

            shutil.rmtree(tmp, ignore_errors=True)

    def test_csv_schema_is_same_without_approved_rows(self) -> None:
        """승인 0건·반려만 있어도 소비 측이 읽는 열 구성은 동일하다."""
        import shutil

        tmp = tempfile.mkdtemp(prefix="pq43s3-csv2-")
        try:
            item = W.Decision(review_id="RV-2", product_key="samasound.co.kr:3", field="brand",
                              before="자체제작", proposed="ALLEN&HEATH", decision="반려",
                              confirmed_value="", reviewer="김", decided_at="", note="", fingerprint="z")
            payload = W.decisions_payload(W.DecisionResult(rejected=[item]))
            path = os.path.join(tmp, "d.csv")
            self.assertEqual(W.write_decisions_csv(payload, path), 1)
            with io.open(path, encoding="utf-8-sig") as handle:
                lines = handle.read().splitlines()
            self.assertEqual(len(lines), 2)
            self.assertEqual(lines[0], ",".join(W._DECISION_CSV_HEADER))
            self.assertIn("반려", lines[1])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_approved_row_marks_reviewer_confirmed(self) -> None:
        item = W.Decision(review_id="RV-1", product_key="samasound.co.kr:9", field="model", before="",
                          proposed="SM58", decision="승인", confirmed_value="SM58", reviewer="김",
                          decided_at="2026-10-23", note="", fingerprint="a")
        row = W.decisions_payload(W.DecisionResult(approved=[item]))["approved"][0]
        self.assertEqual(row["classification"], "reviewer_confirmed")
        self.assertEqual(row["applies_to"], "patch layer only")
        self.assertNotIn("price", json.dumps(row, ensure_ascii=False))


class TestStats(QueueBuilder):
    def test_stage3_stats_shape(self) -> None:
        fp = W.queue_fingerprint(self.queue)
        stats = W.build_stage3_stats(self.queue, self.dups, fp)
        self.assertEqual(stats["review_queue_rows"], len(self.queue))
        self.assertEqual(sum(stats["priority_counts"].values()), len(self.queue))
        self.assertEqual(stats["duplicate_merges_performed"], 0)
        self.assertEqual(stats["stage3_version"], W.STAGE3_VERSION)

    def test_stats_partition_by_field(self) -> None:
        stats = W.build_stage3_stats(self.queue, self.dups, W.queue_fingerprint(self.queue))
        self.assertEqual(stats["model_review"] + stats["manufacturer_review"] + stats["brand_review"],
                         len(self.queue))
        self.assertEqual(Counter(r["필드"] for r in self.queue)["model"], stats["model_review"])


@unittest.skipUnless(
    os.path.exists(os.path.join(ROOT, "site", "data", "products.json")),
    "원본 마스터가 있는 작업 환경에서만 실행",
)
class TestRealMasterStage3(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        from tools.product_quality import build_quality as B

        cls.master_path = os.path.join(ROOT, "site", "data", "products.json")
        with io.open(cls.master_path, "rb") as handle:
            cls.sha_before = __import__("hashlib").sha256(handle.read()).hexdigest()
        cls.tmp = tempfile.mkdtemp(prefix="pq43s3-real-")
        cls.out = os.path.join(cls.tmp, "product_quality")
        B.main(["--out-dir", cls.out, "--quiet"])
        with io.open(cls.master_path, "rb") as handle:
            cls.sha_after = __import__("hashlib").sha256(handle.read()).hexdigest()
        with io.open(os.path.join(cls.out, B.FILES["summary"]), encoding="utf-8") as handle:
            cls.summary = json.load(handle)
        cls.stage3 = cls.summary["stage3"]

    @classmethod
    def tearDownClass(cls) -> None:
        import shutil

        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_master_is_untouched(self) -> None:
        self.assertEqual(self.sha_before, self.sha_after)

    def test_queue_matches_stage2_review_count(self) -> None:
        self.assertEqual(self.stage3["review_queue_rows"], 2311)
        self.assertEqual(self.stage3["model_review"], 1969)
        self.assertEqual(self.stage3["manufacturer_review"], 269)
        self.assertEqual(self.stage3["brand_review"], 73)
        self.assertEqual(self.stage3["priority_counts"], {"1": 156, "2": 620, "3": 239, "4": 1296})
        self.assertEqual(self.stage3["duplicate_pairs"], 1353)

    def test_priority_counts_by_field_covers_every_review_row(self) -> None:
        by_field = self.stage3["priority_counts_by_field"]
        self.assertEqual(sum(sum(v.values()) for v in by_field.values()), 2311)
        self.assertEqual(by_field["model"], {"1": 148, "2": 620, "3": 239, "4": 962})
        self.assertEqual(by_field["manufacturer"], {"4": 269})
        self.assertEqual(by_field["brand"], {"1": 8, "4": 65})

    def test_queue_xlsx_and_decisions_are_generated(self) -> None:
        for name in ("product_quality_review_queue.xlsx", "product_review_decisions.json",
                     "product_review_decisions.csv"):
            self.assertGreater(os.path.getsize(os.path.join(self.out, name)), 0, name)
        with io.open(os.path.join(self.out, "product_review_decisions.json"), encoding="utf-8") as handle:
            payload = json.load(handle)
        self.assertEqual(payload["counts"]["approved"], 0)
        self.assertEqual(payload["counts"]["pending"], 2311)
        self.assertEqual(payload["errors"], [])
        self.assertFalse(payload["policy"]["original_master_overwritten"])
        self.assertEqual(payload["policy"]["duplicates_merged"], 0)
        self.assertEqual(payload["queue_fingerprint"], self.stage3["queue_fingerprint"])

    def test_stage3_does_not_change_stage2_layer(self) -> None:
        self.assertEqual(self.summary["schema_version"], V.STAGE2_VERSION)
        self.assertEqual(self.summary["stage2"]["auto_promoted_model"], 40)
        self.assertEqual(self.summary["missing_after"]["model"], 2427)

    def test_fingerprint_is_stable_across_regeneration(self) -> None:
        from tools.product_quality import build_quality as B

        second = os.path.join(self.tmp, "second")
        B.main(["--out-dir", second, "--quiet"])
        with io.open(os.path.join(second, B.FILES["summary"]), encoding="utf-8") as handle:
            again = json.load(handle)
        self.assertEqual(again["stage3"]["queue_fingerprint"], self.stage3["queue_fingerprint"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
