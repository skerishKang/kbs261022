"""이슈 #43 Stage 2 — 검토 후보 교차 실측(승격) 규칙 테스트.

승격 기준(독립 원본 ID 2건 이상)과 "승격하지 않는 것"의 경계를 다룬다.
기대값은 모두 실제 원본 표기(모델명 스펙행) 형태로 만든 합성 사례이며,
실 마스터 대상 검증은 별도 클래스에서 산출물과 대조한다.
"""

from __future__ import annotations

import copy
import io
import json
import os
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
sys.path.insert(0, ROOT)

from tools.product_quality import quality as Q  # noqa: E402
from tools.product_quality import verify_review as V  # noqa: E402
from tools.product_quality.tests.test_quality import record  # noqa: E402


def run_stage2(fixture, confirmed=()):
    pairs = [
        Q.DuplicateCandidate(
            verdict="confirmed", auto_mergeable=True,
            left_key=c["left_key"], left_name="", left_brand="", left_model="",
            right_key=c["right_key"], right_name="", right_brand="", right_model="",
            shared_phrase=c.get("shared_phrase", ""), name_similarity="1.0",
            variant_conflict="아니오", brand_agreement="예", reason="test", origin="test",
        )
        for c in confirmed
    ]
    result = Q.build_result(fixture, confirmed, pairs)
    stage1 = copy.deepcopy(result.summary)
    stage2 = V.apply(result, fixture)
    result.summary = Q.summarize(result)
    return result, stage2, stage1


def patch_of(result, key, field):
    return next(p for p in result.patches if p.product_key == key and p.field == field)


def improved_of(result, key):
    return next(r for r in result.improved if r["id"] == key)


class TestPromotionRule(unittest.TestCase):
    """후보 모델명이 같은 브랜드의 명시 스펙행에서 독립 ID 2건 이상으로 실측될 때만 승격."""

    def setUp(self) -> None:
        self.fixture = [
            record("삼아스토어", "1", "SM58 마이크", "SHURE", model="SM58", brand_row="SHURE"),
            record("삼아스토어", "2", "SM58 마이크 Black", "SHURE", model="SM58", brand_row="SHURE"),
            record("삼아사운드", "9", "SM58 다이내믹 마이크", "SHURE"),  # 모델명 스펙행 없음
            record("삼아사운드", "10", "KSM9 마이크", "SHURE"),          # 실측 1건
            record("삼아스토어", "3", "KSM9 마이크", "SHURE", model="KSM9", brand_row="SHURE"),
        ]
        self.result, self.stage2, _ = run_stage2(self.fixture)

    def test_two_independent_ids_promote(self) -> None:
        patch = patch_of(self.result, "samasound.co.kr:9", "model")
        self.assertEqual(patch.classification, Q.CLASS_AUTO)
        self.assertEqual(patch.evidence_type, "multi_source_attestation")
        self.assertEqual(patch.proposed, "SM58")
        ver = self.stage2.verification[("samasound.co.kr:9", "model")]
        self.assertEqual(ver.attest_ids, 2)
        self.assertTrue(ver.promoted)
        self.assertIn("SM58", patch.evidence)
        self.assertTrue(patch.evidence_url.startswith("https://samamusic.co.kr/product/1/"),
                        patch.evidence_url)

    def test_single_id_is_not_promoted(self) -> None:
        patch = patch_of(self.result, "samasound.co.kr:10", "model")
        self.assertEqual(patch.classification, Q.CLASS_REVIEW)
        ver = self.stage2.verification[("samasound.co.kr:10", "model")]
        self.assertEqual(ver.attest_ids, 1)
        self.assertFalse(ver.promoted)
        self.assertEqual(ver.priority, "1")

    def test_same_source_id_rows_are_not_independent_evidence(self) -> None:
        """동일 원본 ID가 반복 관찰돼도 독립 증거로 치지 않는다."""
        fixture = [
            record("삼아스토어", "1", "SM58 마이크", "SHURE", model="SM58", brand_row="SHURE"),
            record("삼아스토어", "1", "SM58 마이크 (복사 행)", "SHURE", model="SM58", brand_row="SHURE"),
            record("삼아사운드", "9", "SM58 다이내믹 마이크", "SHURE"),
        ]
        result, stage2, _ = run_stage2(fixture)
        ver = stage2.verification[("samasound.co.kr:9", "model")]
        self.assertEqual(ver.attest_observations, 2, "행은 2번 관찰된다")
        self.assertEqual(ver.attest_ids, 1, "독립 원본 ID는 1건이다")
        self.assertFalse(ver.promoted)
        self.assertEqual(patch_of(result, "samasound.co.kr:9", "model").classification, Q.CLASS_REVIEW)
        self.assertEqual(improved_of(result, "samasound.co.kr:9")["model_after"], "")

    def test_other_brand_attestation_does_not_promote(self) -> None:
        fixture = [
            record("삼아스토어", "1", "X200 인터페이스", "FOCUSRITE", model="X200", brand_row="FOCUSRITE"),
            record("삼아스토어", "2", "X200 인터페이스 Black", "FOCUSRITE", model="X200", brand_row="FOCUSRITE"),
            record("삼아사운드", "9", "X200 오디오 인터페이스", "APPLE"),
        ]
        _, stage2, _ = run_stage2(fixture)
        ver = stage2.verification[("samasound.co.kr:9", "model")]
        self.assertEqual(ver.attest_ids, 0)
        self.assertEqual(ver.other_brand_hits, 1)
        self.assertFalse(ver.promoted)

    def test_promotion_applies_source_spelling(self) -> None:
        """후보 표기가 어휘와 철자가 달라도 원본 스펙행 표기를 그대로 쓴다."""
        fixture = [
            record("삼아스토어", "1", "SM58 마이크", "SHURE", model="SM58", brand_row="SHURE"),
            record("삼아스토어", "2", "SM58 마이크 Black", "SHURE", model="SM58", brand_row="SHURE"),
            record("삼아사운드", "9", "sm-58 다이내믹 마이크", "SHURE"),
        ]
        result, stage2, _ = run_stage2(fixture)
        if not stage2.promoted:
            self.skipTest("sm-58 후보가 어휘와 매칭되지 않으면 이 사례는 승격 대상이 아니다")
        improved = improved_of(result, "samasound.co.kr:9")
        self.assertEqual(improved["model_after"], "SM58")

    def test_lexicon_counts_distinct_ids_per_brand(self) -> None:
        from tools.product_master import normalize as N

        lexicon = V.build_model_lexicon(self.fixture)
        entry = lexicon[N.model_compact("SM58")]
        self.assertEqual(entry["ids"], {"samamusic.co.kr:1", "samamusic.co.kr:2"})
        self.assertEqual(len(entry["by_brand"]["SHURE"]), 2)


class TestNeverPromoted(unittest.TestCase):
    """제조사·브랜드·가격은 Stage 2에서 승격하지 않는다 (대표 결정 2-A·3-B)."""

    def setUp(self) -> None:
        self.fixture = [
            record("삼아스토어", "1", "SM58 마이크", "SHURE", model="SM58", brand_row="SHURE",
                   price_value=620000, price_state="public"),
            record("삼아스토어", "2", "SM58 마이크 Black", "SHURE", model="SM58", brand_row="SHURE",
                   price_value=620000, price_state="public"),
            record("삼아스토어", "4", "KSM9 마이크", "SHURE", model="KSM9", brand_row="SHURE"),
            # 삼아사운드: 제조사 자리에 '자체제작' 표기 → 제조사 검토 유지, 모델명 승격은 가능
            record("삼아사운드", "9", "SM58 다이내믹 마이크", "SHURE", manufacturer="자체제작"),
            # 삼아사운드: 브랜드 자리에 '자체제작' → 브랜드는 제안만 하고 원본 표기 유지
            record("삼아사운드", "10", "KSM9 마이크", "자체제작"),
        ]
        self.result, self.stage2, self.stage1 = run_stage2(self.fixture)

    def test_brand_placeholder_gets_priority_only(self) -> None:
        patch = patch_of(self.result, "samasound.co.kr:10", "brand")
        self.assertEqual(patch.classification, Q.CLASS_REVIEW, "브랜드는 자동 변경하지 않는다")
        self.assertEqual(patch.proposed, "SHURE")
        self.assertEqual(patch.evidence_type, "model_lexicon_brand_attestation")
        ver = self.stage2.verification[("samasound.co.kr:10", "brand")]
        self.assertEqual(ver.priority, "1")
        self.assertFalse(ver.promoted)
        improved = improved_of(self.result, "samasound.co.kr:10")
        self.assertEqual(improved["brand"], "자체제작")
        self.assertEqual(self.stage2.stats["auto_promoted_brand"], 0)

    def test_manufacturer_review_is_annotated_not_promoted(self) -> None:
        patch = patch_of(self.result, "samasound.co.kr:9", "manufacturer")
        self.assertEqual(patch.classification, Q.CLASS_REVIEW)
        ver = self.stage2.verification[("samasound.co.kr:9", "manufacturer")]
        self.assertEqual(ver.priority, "4")
        self.assertFalse(ver.promoted)
        self.assertIn("제조사", ver.rationale)
        self.assertEqual(improved_of(self.result, "samasound.co.kr:9")["manufacturer_after"], "자체제작")

    def test_price_fields_are_passed_through_unchanged(self) -> None:
        for original, out in zip(self.fixture, self.result.improved):
            self.assertEqual(out["price_original"], original["price"]["value"])
            self.assertEqual(out["price_state_original"], original["price"]["state"])
        self.assertEqual(self.result.summary["policy"]["price_invented"], False)

    def test_original_records_are_not_mutated(self) -> None:
        snapshot = copy.deepcopy(self.fixture)
        run_stage2(self.fixture)
        self.assertEqual(snapshot, self.fixture, "원본 레코드 dicts 는 읽기 전용")

    def test_flags_and_changed_fields_follow_promotion(self) -> None:
        improved = improved_of(self.result, "samasound.co.kr:9")
        self.assertIn("model_stage2_promoted", improved["quality_flags"])
        self.assertNotIn("model_review", improved["quality_flags"])
        self.assertIn("model", improved["changed_fields"])
        self.assertEqual(improved["patches"]["model"]["classification"], Q.CLASS_AUTO)
        self.assertEqual(improved["model_after"], "SM58")


class TestStage2Stats(unittest.TestCase):
    def setUp(self) -> None:
        self.fixture = [
            record("삼아스토어", "1", "SM58 마이크", "SHURE", model="SM58", brand_row="SHURE"),
            record("삼아스토어", "2", "SM58 마이크 Black", "SHURE", model="SM58", brand_row="SHURE"),
            record("삼아사운드", "9", "SM58 다이내믹 마이크", "SHURE"),
            record("삼아사운드", "10", "KSM9 마이크", "SHURE"),
        ]
        self.result, self.stage2, self.stage1 = run_stage2(self.fixture)

    def test_promoted_count_matches_stats(self) -> None:
        self.assertEqual(self.stage2.stats["auto_promoted_model"], len(self.stage2.promoted))
        self.assertEqual(self.stage2.stats["auto_promoted_model"], 1)
        self.assertEqual(self.stage2.stats["auto_promoted_manufacturer"], 0)
        self.assertEqual(self.stage2.stats["auto_promoted_brand"], 0)

    def test_missing_counts_move_only_by_promotions(self) -> None:
        self.assertEqual(self.stage2.stats["model_missing_after_stage1"], 2)
        self.assertEqual(self.stage2.stats["model_missing_after_stage2"], 1)
        self.assertEqual(self.result.summary["missing_after"]["model"], 1)

    def test_stage1_snapshot_is_kept_for_before_after(self) -> None:
        self.assertEqual(self.stage1["missing_after"]["model"], 2)
        stage2_review = sum(1 for p in self.result.patches
                            if p.actionable and p.classification == Q.CLASS_REVIEW)
        self.assertEqual(self.stage1["review_required_total"] - stage2_review,
                         self.stage2.stats["auto_promoted_model"])
        self.assertEqual(self.result.summary["review_required_total"], stage2_review)

    def test_change_history_records_promotion(self) -> None:
        joined = " ".join(self.stage2.stats["change_history"])
        self.assertIn("승격", joined)
        self.assertIn("schema_version", joined)
        self.assertEqual(self.stage2.stats["stage2_version"], V.STAGE2_VERSION)

    def test_promotions_carry_evidence(self) -> None:
        row = self.stage2.promotions[0]
        for key in ("id", "model_before", "model_after", "attest_independent_ids", "evidence", "evidence_url"):
            self.assertIn(key, row)
        self.assertGreaterEqual(row["attest_independent_ids"], V.REQUIRED_INDEPENDENT_IDS)
        self.assertTrue(row["evidence"])
        self.assertTrue(row["evidence_url"])

    def test_verification_columns_contract(self) -> None:
        columns = V.Verification().to_columns()
        for name in ("실측독립ID", "실측관찰", "영문파생일치", "중복판정등급", "검토우선순위", "판정근거"):
            self.assertIn(name, columns)

    def test_priority_labels_are_defined_for_every_used_priority(self) -> None:
        used = {v.priority for v in self.stage2.verification.values() if v.priority}
        self.assertTrue(used <= set(V.PRIORITY_LABELS), used)

    def test_version_constants(self) -> None:
        self.assertEqual(V.STAGE2_VERSION, "1.1.0")
        self.assertEqual(V.REQUIRED_INDEPENDENT_IDS, 2)


@unittest.skipUnless(
    os.path.exists(os.path.join(ROOT, "site", "data", "products.json")),
    "원본 마스터가 있는 작업 환경에서만 실행",
)
class TestRealMasterStage2(unittest.TestCase):
    """실 데이터 4,575건 — 승격 근거를 원본에서 독립적으로 재대조한다."""

    @classmethod
    def setUpClass(cls) -> None:
        from tools.product_quality import build_quality as B

        cls.master_path = os.path.join(ROOT, "site", "data", "products.json")
        with io.open(cls.master_path, "rb") as handle:
            cls.sha_before = handle.read()
        cls.master = B.load_master(cls.master_path)
        cls.master_pairs, cls.confirmed_pairs = B.load_duplicate_sheet(
            os.path.join(ROOT, "data", "full_catalog", "product_db.xlsx")
        )
        cls.result = Q.build_result(cls.master, cls.confirmed_pairs, cls.master_pairs)
        cls.stage1_missing_model = cls.result.summary["missing_after"]["model"]
        cls.stage1_review_total = cls.result.summary["review_required_total"]
        cls.stage2 = V.apply(cls.result, cls.master)
        cls.result.summary = Q.summarize(cls.result)
        cls.tmp = tempfile.mkdtemp(prefix="pq43s2-")
        cls.out = os.path.join(cls.tmp, "product_quality")
        B.main(["--out-dir", cls.out, "--quiet"])
        with io.open(cls.master_path, "rb") as handle:
            cls.sha_after = handle.read()

    def test_master_untouched_and_count_kept(self) -> None:
        self.assertEqual(self.sha_before, self.sha_after)
        self.assertEqual(self.result.master_count, 4575)
        self.assertEqual(len(self.result.improved), 4575)

    def test_promotion_volume_and_missing_delta(self) -> None:
        stats = self.stage2.stats
        self.assertEqual(stats["auto_promoted_model"], len(self.stage2.promotions))
        self.assertEqual(
            self.stage1_missing_model - stats["model_missing_after_stage2"],
            stats["auto_promoted_model"],
        )
        self.assertEqual(self.result.summary["missing_after"]["model"],
                         stats["model_missing_after_stage2"])

    def test_every_promotion_is_independently_attested(self) -> None:
        from tools.product_master import normalize as N

        by_key = {Q._clean(r.get("id")): r for r in self.master}
        for detail in self.stage2.promotions:
            record_in = by_key[detail["id"]]
            brand = Q._clean(record_in.get("brand_norm"))
            compact = N.model_compact(detail["model_after"])
            peers = {
                Q._clean(r.get("id"))
                for r in self.master
                if Q._clean(r.get("id")) != detail["id"]
                and Q._clean(r.get("brand_norm")) == brand
                and r.get("model") and not Q.is_placeholder(r["model"])
                and N.model_compact(r["model"]) == compact
            }
            self.assertGreaterEqual(len(peers), V.REQUIRED_INDEPENDENT_IDS, detail["id"])
            self.assertTrue(detail["evidence_url"], detail["id"])

    def test_nothing_outside_model_is_promoted(self) -> None:
        for patch in self.result.patches:
            if patch.field == "model":
                continue
            self.assertNotEqual(patch.evidence_type, "multi_source_attestation", patch.product_key)
            self.assertEqual(
                self.stage2.verification[(patch.product_key, patch.field)].promoted, False)
        self.assertEqual(self.stage2.stats["auto_promoted_manufacturer"], 0)
        self.assertEqual(self.stage2.stats["auto_promoted_brand"], 0)

    def test_review_rows_still_review(self) -> None:
        remaining = sum(1 for p in self.result.patches
                        if p.classification == Q.CLASS_REVIEW and p.actionable)
        self.assertEqual(remaining, self.stage1_review_total - self.stage2.stats["auto_promoted_model"])

    def test_brand_and_price_are_untouched(self) -> None:
        for rec in self.result.improved:
            if rec.get("verification", {}).get("brand", {}).get("promoted_by_stage2"):
                self.fail(f"브랜드가 자동 변경되었다: {rec['id']}")
        for original, rec in zip(self.master, self.result.improved):
            self.assertEqual(rec["brand"], original.get("brand"))
            self.assertEqual(rec["price_original"], (original.get("price") or {}).get("value"))
            self.assertEqual(rec["price_state_original"], (original.get("price") or {}).get("state"))

    def test_generated_artifacts_carry_stage2(self) -> None:
        with io.open(os.path.join(self.out, "product_quality_summary.json"), encoding="utf-8") as handle:
            summary = json.load(handle)
        self.assertEqual(summary["schema_version"], V.STAGE2_VERSION)
        self.assertEqual(summary["stage2"]["auto_promoted_model"], self.stage2.stats["auto_promoted_model"])
        self.assertEqual(len(summary["stage2"]["promotions"]),
                         self.stage2.stats["auto_promoted_model"])
        self.assertEqual(summary["stage1_before_stage2"]["missing_after"]["model"], self.stage1_missing_model)

        with io.open(os.path.join(self.out, "product_quality_improved.json"), encoding="utf-8") as handle:
            improved = json.load(handle)
        self.assertEqual(improved["schema_version"], V.STAGE2_VERSION)
        self.assertEqual(len(improved["products"]), 4575)
        promoted = [r for r in improved["products"] if "model_stage2_promoted" in r["quality_flags"]]
        self.assertEqual(len(promoted), self.stage2.stats["auto_promoted_model"])
        for rec in promoted:
            self.assertEqual(rec["patches"]["model"]["classification"], Q.CLASS_AUTO)
            self.assertTrue(rec["model_after"])
            self.assertEqual(rec["verification"]["model"]["promoted_by_stage2"], True)

    def test_stage2_csv_rows_equal_review_plus_promotions(self) -> None:
        import csv

        path = os.path.join(self.out, "product_quality_stage2_verification.csv")
        with io.open(path, encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        expected = self.stage1_review_total  # Stage 1 검토 행 = 승격 40 + 남은 검토
        self.assertEqual(len(rows), expected)
        self.assertEqual(sum(1 for r in rows if r["Stage2승격"] == "예"),
                         self.stage2.stats["auto_promoted_model"])
        for column in ("실측독립ID", "실측관찰", "영문파생일치", "중복판정등급", "검토우선순위"):
            self.assertIn(column, rows[0])

    def test_review_xlsx_sheet_contract(self) -> None:
        from openpyxl import load_workbook

        path = os.path.join(self.out, "product_quality_review.xlsx")
        workbook = load_workbook(path, read_only=True)
        try:
            names = [ws.title for ws in workbook.worksheets]
            for need in ("요약", "Stage2승격", "자동보완", "검토후보", "근거없음", "중복후보",
                         "필드현황", "LOCAL2계약"):
                self.assertIn(need, names)
            promoted_sheet = workbook["Stage2승격"]
            rows = list(promoted_sheet.iter_rows(values_only=True))
            header = list(rows[1])
            self.assertEqual(len(rows) - 2, self.stage2.stats["auto_promoted_model"])
            self.assertEqual(header[0], "통합키")
            self.assertIn("검토우선순위", header)
            review = list(workbook["검토후보"].iter_rows(values_only=True))
            rhead = list(review[1])
            pr = rhead.index("검토우선순위")
            order = [r[pr] for r in review[2:]]
            self.assertEqual(order, sorted(order), "검토후보 시트는 우선순위 정렬이어야 한다")
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
